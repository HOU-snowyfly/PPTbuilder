"""Adult editorial stories and image-free magazine layouts."""
from pathlib import Path
from pptx.util import Inches
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from visuals import Canvas, INK, WHITE, CORAL
NAVY='153E75'; PALE='EEF4FA'; ROSE='FFF0F1'; PAPER='F7F6F3'
DIR=Path(__file__).parent/'assets'/'editorial'
EDITORIAL_ARTWORKS=[('key_stories.png','小周、阿明与小林的三个重点故事','2026-10-13'),('support_scenes.png','晓岚的课堂故事与三个轻量场景','2026-10-13')]
EDITORIAL_USAGE=[(n,'key_stories.png',f'R{r+1}C{j+1}',desc) for n,r,descs in [(6,0,['症状仍在','室友递第二盒','展开两份说明书']),(18,1,['白天服药与聚餐邀请','室友劝酒','选白水并查说明书']),(27,2,['已服药','想加第二盒','准备聚餐'])] for j,desc in enumerate(descs)]
EDITORIAL_USAGE += [(4,'support_scenes.png','R1C1','晓岚鼻部症状缓解'),(4,'support_scenes.png','R1C2','晓岚课堂困倦'),(8,'support_scenes.png','R2C1','小周核对记录'),(22,'support_scenes.png','R2C2','手部核对完整药板与说明书'),(23,'support_scenes.png','R3C1','成年学生咨询药师')]
BOUNDS={
 'key_stories.png':([(0,503),(516,1020),(1032,1536)],[(0,330),(339,664),(673,1024)]),
 'support_scenes.png':([(0,760),(776,1536)],[(0,330),(339,663),(672,1024)])}

def scene(c,file,row,col,x,y,w,h):
    cols,rows=BOUNDS[file];l,r=cols[col];t,b=rows[row];sw,sh=r-l,b-t;ratio=w/h
    if sw/sh>ratio:
        cut=(sw-sh*ratio)/2;l+=cut;r-=cut
    else:
        cut=sh-sw/ratio;t+=cut*.10;b-=cut*.90
    pic=c.s.shapes.add_picture(str(DIR/file),Inches(x),Inches(y),Inches(w),Inches(h))
    pic.crop_left=l/1536;pic.crop_right=1-r/1536;pic.crop_top=t/1024;pic.crop_bottom=1-b/1024
    pic.name=f"EDITORIAL_{'KEY' if file.startswith('key') else 'SUPPORT'}_R{row+1}_C{col+1}"
    pic._element.xpath('.//p:cNvPr')[0].set('descr','成熟出版物编辑插画；不同案例更换人物，同一连续故事保持一致；抽象药盒为情境道具')
    return pic

def white(c):
    c.s.background.fill.solid();c.s.background.fill.fore_color.rgb=RGBColor.from_string(WHITE)

def cover(c,d):
    white(c);c.box(.65,.46,.08,.28,CORAL)
    c.text('苏州大学 · 校园用药课堂',.87,.40,7.0,.44,18,NAVY,True)
    c.text('合理用药',.80,1.52,6.22,1.45,88,NAVY,True)
    c.box(.88,3.07,1.15,.08,CORAL)
    c.text('给宿舍药箱上堂课',.88,3.38,6.04,.62,31,NAVY,True)
    c.box(7.20,1.65,5.22,2.39,PAPER)
    c.text('从一次“吃什么药”\n到一个有依据的决定。',7.48,2.01,4.66,1.34,30,NAVY,True)
    for j,t in enumerate(['看说明书','核对成分','问清疑点','留好记录']):
        x=.88+j*3.02;c.box(x,4.88,2.80,1.03,ROSE if j==1 else PALE)
        c.text(t,x+.18,5.18,2.44,.43,24,NAVY,True,True)
    c.text(d['items'][2],.90,6.92,6.62,.35,17,NAVY)
    c.text('35分钟 · 双人展示',9.01,6.92,3.42,.35,17,NAVY,False,True)

def agenda(c,d):
    for j,item in enumerate(d['items']):
        title,time=item.split('｜');x=.85+(j%2)*5.98;y=2.13+(j//2)*2.05
        c.box(x,y,5.63,1.80,PALE if j%2==0 else ROSE)
        c.text(title[:2],x+.27,y+.22,1.08,.55,32,CORAL,True)
        c.text(title[3:],x+1.61,y+.20,3.64,.59,30,NAVY,True)
        c.text(time,x+1.61,y+1.04,3.64,.42,23,INK)

def quiz(c,d):
    c.box(10.69,.37,1.70,.41,NAVY)
    c.text('可多选' if d.get('question_mode')=='多选' else '单选',10.80,.40,1.48,.34,16,WHITE,True,True)
    for j,item in enumerate(d['items']):
        y=2.15+j*1.34;c.box(.85,y,11.60,1.05,PALE)
        c.box(1.11,y+.26,.59,.55,NAVY)
        c.text(item[0],1.11,y+.31,.59,.44,22,WHITE,True,True)
        c.text(item[2:],2.00,y+.21,10.09,.65,26,INK)

def answer(c,d):
    result='A＋C' if d['n']==10 else 'A＋B'
    c.box(.85,2.13,3.08,4.03,NAVY);c.text(result,1.08,2.72,2.62,1.00,52,WHITE,True,True)
    c.text('把信息补齐\n再决定下一步',1.15,4.24,2.48,1.12,24,WHITE,True,True)
    for j,item in enumerate(d['items']):
        y=2.13+j*1.34;c.box(4.28,y,8.17,1.20,ROSE if j==2 else PALE)
        if '｜' in item:
            title,body=item.split('｜');c.text(title,4.56,y+.15,7.61,.37,25,NAVY,True);c.text(body,4.56,y+.65,7.61,.37,22,INK)
        else:c.text(item,4.56,y+.22,7.61,.75,25,NAVY,True)

def three_story(c,d):
    row={6:0,18:1,27:2}[d['n']]
    for j,item in enumerate(d['items']):
        x=.82+j*4.03;scene(c,'key_stories.png',row,j,x,2.14,3.79,2.52)
        c.caption(item,x,4.83,3.79,1.12,21)

def timeline(c,d):
    scene(c,'support_scenes.png',1,0,.85,2.16,5.63,2.50)
    c.box(.85,4.94,5.63,1.18,ROSE)
    c.text('还缺什么信息？',1.11,5.09,5.11,.34,24,NAVY,True)
    c.text('成分、规格与今天的已服情况',1.11,5.57,5.11,.35,22,INK)
    for j,item in enumerate(d['items'][:2]):
        title,body=item.split('｜');y=2.16+j*2.02;c.box(6.83,y,5.63,1.91,PALE)
        c.text(title,7.10,y+.20,5.09,.52,32,NAVY,True)
        c.text(body,7.10,y+1.05,5.09,.58,24,INK)

def allergy(c,d):
    for j,x in enumerate([.85,6.83]):
        scene(c,'support_scenes.png',0,j,x,2.15,5.63,2.70)
        c.caption(['鼻痒、流涕缓解了。','到了课堂，却开始犯困。'][j],x,5.04,5.63,1.02,24)

def antibiotics(c,d):
    for j,item in enumerate(d['items']):
        title,body=item.split('｜');x=.85+j*4.04
        c.shape(MSO_SHAPE.RECTANGLE,x,2.16,3.80,3.92,ROSE if j==1 else PALE)
        c.text(title,x+.28,2.53,3.24,.87,29,NAVY,True)
        body=['感冒症状，\n不自动等于要用抗生素','相似症状，\n不代表同一病因','疗效或反应有疑问，\n及时咨询'][j]
        c.text(body,x+.28,3.83,3.24,1.64,23,INK)

def framework(c,d):
    for j,item in enumerate(d['items']):
        key,body=item.split('｜');x=.85+(j%2)*5.98;y=2.13+(j//2)*2.05
        c.box(x,y,5.63,1.80,ROSE if j%2 else PALE)
        c.text(key,x+.26,y+.30,1.03,.76,45,NAVY,True)
        body=['症状、说明书、\n就医信号','成分、用法、禁忌\n与相互作用','带齐信息，\n问医生或药师','药名、时间、\n用量与反应'][j]
        c.text(body,x+1.54,y+.30,3.81,1.20,24,INK)

def leaflet(c,d):
    for j,item in enumerate(d['items']):
        title,body=item.split('｜');x=.85+(j%2)*5.98;y=2.13+(j//2)*1.35
        c.shape(MSO_SHAPE.RECTANGLE,x,y,5.63,1.19,PALE)
        c.text(title,x+.26,y+.15,5.11,.35,23,NAVY,True)
        c.text(body,x+.26,y+.65,5.11,.37,23,INK)

def operation(c,d):
    scene(c,'support_scenes.png',1,1,.85,2.29,5.63,2.50)
    c.box(.85,5.07,5.63,1.05,PAPER);c.text('先核具体要求，\n不凭感觉改变使用方式。',1.13,5.21,5.07,.75,24,NAVY,True)
    for j,item in enumerate(d['items']):
        title,body=item.split('｜');y=2.13+j*1.36;c.box(6.83,y,5.63,1.22,ROSE if j==2 else PALE)
        c.text(title,7.10,y+.15,5.09,.36,26,NAVY,True);c.text(body,7.10,y+.65,5.09,.40,22,INK)

def consult(c,d):
    scene(c,'support_scenes.png',2,0,.85,2.29,5.63,2.50)
    c.box(.85,5.07,5.63,1.05,PAPER);c.text('药盒、说明书和已服记录，\n一起带来。',1.13,5.21,5.07,.75,24,NAVY,True)
    for j,item in enumerate(d['items']):
        y=2.13+j*1.36;c.box(6.83,y,5.63,1.22,PALE)
        c.text(item,7.10,y+.20,5.09,.83,25,NAVY,True)

def urgent(c,d):
    c.box(.85,2.15,11.60,1.45,CORAL)
    c.text('120',1.13,2.35,2.30,1.01,64,WHITE,True)
    c.text('呼吸困难、喉舌肿胀或意识异常',4.09,2.39,7.93,.45,28,WHITE,True)
    c.text('严重急症：立即寻求急救',4.09,3.03,7.93,.34,24,WHITE,True)
    for j,item in enumerate(d['items'][1:]):
        title,body=item.split('｜');x=.85+j*5.98;c.box(x,3.98,5.63,2.10,PALE)
        c.text(title,x+.27,4.30,5.09,.66,27,NAVY,True);c.text(body,x+.27,5.29,5.09,.45,25,INK)

def kit(c,d):
    for j,item in enumerate(d['items'][:2]):
        title,body=item.split('｜');x=.85+j*5.98;c.box(x,2.14,5.63,1.65,PALE)
        c.text(title,x+.27,2.39,5.09,.43,28,NAVY,True);c.text(body,x+.27,3.11,5.09,.42,23,INK)
    c.shape(MSO_SHAPE.RECTANGLE,.85,4.08,11.60,2.09,ROSE);c.text('留好用药记录',1.12,4.31,11.06,.42,28,NAVY,True)
    for j,t in enumerate(['药名','时间','用量','反应']):
        x=1.12+j*2.81;c.text(t,x,5.05,2.29,.43,25,NAVY,True);c.line(x,5.79,x+2.26,5.79,NAVY,1.0)

def recap(c,d):
    for j,item in enumerate(d['items']):
        situation,act=item.split(' → ');y=2.13+j*1.02;c.box(.85,y,11.60,.91,ROSE if j%2 else PALE)
        c.text(situation,1.13,y+.22,3.25,.47,23,NAVY,True);c.arrow(4.79,y+.32,.45,CORAL);c.text(act,5.59,y+.22,6.55,.47,23,NAVY,True)

def closing(c,d):
    white(c);c.box(.65,.46,.08,.28,CORAL);c.text('合理用药',.87,.40,7.0,.44,18,NAVY,True)
    c.text('药可以常备，',.86,1.20,11.62,.71,38,NAVY,True)
    c.text('判断不能省略。',.86,2.12,11.62,.99,57,NAVY,True)
    for j,item in enumerate(d['items']):
        y=3.58+j*.91;c.shape(MSO_SHAPE.RECTANGLE,.85,y,11.60,.76,ROSE if j==1 else PALE)
        c.text(['看','核·记','问'][j],1.03,y+.18,1.20,.40,22,CORAL,True,True);c.text(item,2.56,y+.17,9.49,.43,26,NAVY,True)
    c.text('谢谢大家',9.57,6.84,2.86,.48,25,NAVY,True,True)

DRAWERS={1:cover,5:agenda,2:quiz,9:quiz,19:quiz,28:quiz,10:answer,29:answer,6:three_story,18:three_story,27:three_story,8:timeline,4:allergy,14:antibiotics,20:framework,21:leaflet,22:operation,23:consult,24:urgent,25:kit,26:recap,30:closing}
def draw_editorial_slide(slide,d):
    drawer=DRAWERS.get(d['n'])
    if drawer is None:return False
    drawer(Canvas(slide),d);return True
