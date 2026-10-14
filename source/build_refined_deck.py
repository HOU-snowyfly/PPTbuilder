"""Build the approved refined edition while keeping the original delivery intact."""
from __future__ import annotations

import hashlib
import io
import json
from collections import Counter
from copy import deepcopy
from pathlib import Path

from docx import Document
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

import build_refined_editorial_samples as approved
from build_editorial_texture_samples import clone_shapes, gradient, nav, footer
from build_global_style_samples import box, canonical, original_texts, picture, rgb, text


ROOT = Path(__file__).resolve().parents[1]
ORIGINAL = ROOT / "交付文件/合理用药_双人35分钟.pptx"
MANUSCRIPT = ROOT / "交付文件/合理用药_双人逐页讲稿.docx"
OUT = ROOT / "交付文件/精致编辑版"
WORK = ROOT / "work/refined-deck"
PPT_NAME = "合理用药_精致编辑版_双人35分钟.pptx"
DOC_NAME = "合理用药_精致编辑版_双人逐页讲稿.docx"
T = approved.THEME
COLOR_MAP = {"203354": T["ink"], "E3464F": T["accent"], "EB5969": T["accent"],
             "F0F4FA": "FCFAF5", "FFEEF0": "F7ECE7",
             "F0F8F6": "F5F7F2", "F5F1FA": "F6F2F4", "FFFFFF": "FCFBF7"}


def chrome(presentation, original):
    slide, v = approved.base(presentation, original)
    nav(slide, v[0])
    text(slide, v[1], .65, 1.01, 12.02, .82, 32 if len(v[1]) < 25 else 30, bold=True)
    text(slide, v[2], .65, 6.59, 12.02, .53, 20)
    footer(slide, v[3])
    return slide, v[4:]


def card(slide, x, y, w, h, fill="FCFAF5"):
    return box(slide, x, y, w, h, fill)


def cover(presentation, original):
    s, v = approved.base(presentation, original)
    gradient(s, 0, 0, 7.08, 7.5, "203D4D", "315A6A")
    text(s, v[0], .65, .42, 6.02, .4, 14, color=T["paper"])
    text(s, v[1], .63, 1.65, 6.0, 1.45, 82, color=T["paper"], bold=True)
    text(s, v[2], .69, 3.51, 6.0, .58, 28, color=T["paper"])
    quote = text(s, v[3], 7.72, 2.00, 4.88, 1.56, 27)
    quote.text_frame.paragraphs[1].font.bold = True
    card(s, .65, 5.24, 12.02, 1.02, "F8F5EE")
    for index, value in enumerate(v[4:8]):
        x = .92 + index * 3.0
        text(s, value, x, 5.54, 2.43, .4, 23)
        if index:
            box(s, x - .2, 5.54, .01, .39, T["rule"])
    text(s, v[8], .69, 6.98, 6.0, .3, 13, color=T["paper"])
    text(s, v[9], 8.0, 6.98, 4.65, .3, 13, align=PP_ALIGN.RIGHT)
    return s


def route(presentation, original):
    s, v = chrome(presentation, original)
    assert len(v) == 12
    for index in range(4):
        y = 2.05 + index * 1.06
        card(s, .65, y, 12.02, .9)
        text(s, v[index*3], .92, y + .14, 1.0, .52, 31, color=T["accent"])
        text(s, v[index*3+1], 2.30, y + .19, 6.6, .43, 25, bold=True)
        text(s, v[index*3+2], 10.03, y + .23, 2.22, .37, 21, align=PP_ALIGN.RIGHT)
    return s


def quiz(presentation, original):
    s, v = chrome(presentation, original)
    assert len(v) == 7
    text(s, v[0], 10.76, .23, 1.9, .31, 13, align=PP_ALIGN.RIGHT)
    for index in range(3):
        y = 2.17 + index * 1.32
        card(s, .65, y, 12.02, 1.02)
        text(s, v[index*2+1], .94, y + .27, .64, .45, 24, color=T["ink"])
        box(s, 1.78, y + .21, .012, .60, T["rule"])
        text(s, v[index*2+2], 2.05, y + .24, 10.28, .52, 24)
    return s


def principles(presentation, original):
    s, v = chrome(presentation, original)
    assert len(v) == 12
    for index, (x, y) in enumerate(((.65, 2.13), (6.94, 2.13),
                                  (.65, 4.20), (6.94, 4.20))):
        card(s, x, y, 5.73, 1.78)
        text(s, v[index*3], x+.26, y+.29, .58, .29, 14, color=T["accent"])
        text(s, v[index*3+1], x+1.04, y+.20, 4.36, .53, 27, bold=True)
        text(s, v[index*3+2], x+.26, y+1.02, 5.20, .48, 22)
    return s


def clone_body(presentation, original):
    """Keep ingredient and scientific topology, image pixels, crops and labels intact."""
    s, _ = chrome(presentation, original)
    for source in list(original.shapes)[6:]:
        element = deepcopy(source._element)
        element.find(".//"+qn("p:cNvPr")).set("id", str(s.shapes._next_shape_id))
        if source.shape_type == 13:
            _, rid = s.part.get_or_add_image_part(io.BytesIO(source.image.blob))
            element.find(".//"+qn("a:blip")).set(qn("r:embed"), rid)
        s.shapes._spTree.insert_element_before(element, "p:extLst")
        fresh = s.shapes[-1]
        for color in fresh._element.iter(qn("a:srgbClr")):
            old = color.get("val")
            if old in COLOR_MAP:
                color.set("val", COLOR_MAP[old])
        if fresh.has_text_frame:
            for paragraph in fresh.text_frame.paragraphs:
                for run in paragraph.runs:
                    size = run.font.size.pt if run.font.size else 20
                    run.font.name = T["font"]
                    run.font.bold = (size >= 27 or fresh.text in
                                     ("药盒 A", "药盒 B", "对乙酰氨基酚"))
        for effect in fresh._element.iter(qn("a:effectRef")):
            effect.set("idx", "0")
    return s


def timeline(presentation, original):
    s, v = chrome(presentation, original)
    panel = next(p for p in original.shapes if p.shape_type == 13)
    picture(s, panel, 0, 2.19, 6.56, 2.95)
    text(s, v[0], .65, 5.35, 5.7, .4, 22, bold=True)
    text(s, v[1], .65, 5.90, 5.7, .42, 22)
    for index in range(2):
        y = 2.18 + index * 2.03
        card(s, 7.04, y, 5.63, 1.74)
        text(s, v[2+index*2], 7.33, y+.23, 5.05, .52, 30, color=T["accent"])
        text(s, v[3+index*2], 7.33, y+1.04, 5.05, .44, 22)
    return s


def answer(presentation, original):
    s, v = chrome(presentation, original)
    gradient(s, .65, 2.16, 3.26, 4.12, "203F50", "335C6D")
    text(s, v[0], .92, 2.75, 2.72, .93, 51, color="EEB9AF")
    text(s, v[1], .96, 4.38, 2.65, 1.00, 22, color=T["paper"])
    paired = len(v) == 8
    for index in range(3):
        y = 2.16 + index * 1.37
        card(s, 4.29, y, 8.38, 1.17)
        if paired:
            text(s, v[2+index*2], 4.55, y+.16, 7.86, .35, 22, bold=True)
            text(s, v[3+index*2], 4.55, y+.67, 7.86, .33, 21)
        else:
            text(s, v[2+index], 4.55, y+.30, 7.86, .59, 23)
    return s


def three_actions(presentation, original):
    s, v = chrome(presentation, original)
    for index in range(3):
        x = .65 + index * 4.19
        card(s, x, 2.2, 3.63, 3.91)
        text(s, v[index*2], x+.25, 2.56, 3.13, .84, 25, bold=True)
        text(s, v[index*2+1], x+.25, 3.93, 3.13, 1.55, 22)
    return s


def two_scenes(presentation, original):
    s, v = chrome(presentation, original)
    panels = [p for p in original.shapes if p.shape_type == 13]
    for index, x in enumerate((.65, 6.94)):
        card(s, x, 2.17, 5.73, 3.93)
        picture(s, panels[index], x, 2.17, 5.73, 2.83)
        text(s, v[index], x+.24, 5.28, 5.25, .53, 22)
    return s


def three_scenes(presentation, original):
    s, v = chrome(presentation, original)
    panels = [p for p in original.shapes if p.shape_type == 13]
    picture(s, panels[0], .65, 2.13, 5.42, 3.67)
    text(s, v[0], .65, 5.95, 5.42, .46, 21)
    for index, x in enumerate((6.70, 9.81), 1):
        card(s, x, 2.36, 2.85, 3.54)
        picture(s, panels[index], x, 2.36, 2.85, 1.90)
        text(s, v[index].replace("，", "，\n", 1), x+.16, 4.56, 2.53, 1.12, 20)
    return s


def leaflet(presentation, original):
    s, v = chrome(presentation, original)
    for index in range(6):
        x = .65 if index % 2 == 0 else 6.94
        y = 2.08 + (index//2)*1.40
        card(s, x, y, 5.73, 1.17)
        text(s, v[index*2], x+.26, y+.16, 5.2, .34, 20, bold=True)
        text(s, v[index*2+1], x+.26, y+.69, 5.2, .35, 22)
    return s


def single_scene(presentation, original):
    s, v = chrome(presentation, original)
    panel = next(p for p in original.shapes if p.shape_type == 13)
    picture(s, panel, 0, 2.17, 6.60, 3.02)
    text(s, v[0], .65, 5.42, 5.73, .91, 22)
    paired = len(v) == 7
    for index in range(3):
        y = 2.17 + index * 1.37
        card(s, 7.03, y, 5.63, 1.14)
        if paired:
            text(s, v[1+index*2], 7.30, y+.13, 5.09, .33, 22, bold=True)
            text(s, v[2+index*2], 7.30, y+.66, 5.09, .33, 20)
        else:
            text(s, v[1+index], 7.30, y+.23, 5.09, .67, 21)
    return s


def urgent(presentation, original):
    s, v = chrome(presentation, original)
    card(s, .65, 2.14, 12.02, 1.69, "F7ECE7")
    text(s, v[0], .94, 2.46, 2.28, .93, 62, color=T["accent"], bold=True)
    text(s, v[1], 4.08, 2.50, 8.14, .43, 24, bold=True)
    text(s, v[2], 4.08, 3.13, 8.14, .33, 21)
    for index, x in enumerate((.65, 6.94)):
        card(s, x, 4.23, 5.73, 1.88)
        text(s, v[3+index*2], x+.27, 4.54, 5.19, .64, 22, bold=True)
        text(s, v[4+index*2], x+.27, 5.47, 5.19, .36, 21)
    return s


def kit(presentation, original):
    s, v = chrome(presentation, original)
    for index, x in enumerate((.65, 6.94)):
        card(s, x, 2.16, 5.73, 1.78)
        text(s, v[index*2], x+.26, 2.46, 5.20, .41, 24, bold=True)
        text(s, v[index*2+1], x+.26, 3.20, 5.20, .39, 22)
    gradient(s, .65, 4.30, 12.02, 1.93, "203F50", "335C6D", angle=0)
    text(s, v[4], .92, 4.59, 11.48, .4, 23, color=T["paper"])
    for index, value in enumerate(v[5:]):
        text(s, value, .92+index*2.9, 5.39, 2.50, .39, 23, color="EEB9AF")
    return s


def recap(presentation, original):
    s, v = chrome(presentation, original)
    for index in range(4):
        y = 2.14 + index*1.04
        card(s, .65, y, 12.02, .86)
        text(s, v[index*2], .93, y+.21, 3.63, .39, 21)
        arrow = s.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(4.63), Inches(y+.29),
                                   Inches(.25), Inches(.23))
        arrow.fill.solid(); arrow.fill.fore_color.rgb = rgb(T["accent"])
        arrow.line.fill.background()
        arrow._element.spPr.append(OxmlElement("a:effectLst"))
        for effect in arrow._element.iter(qn("a:effectRef")): effect.set("idx", "0")
        text(s, v[index*2+1], 5.15, y+.21, 7.20, .39, 22)
    return s


def closing(presentation, original):
    s, v = approved.base(presentation, original)
    gradient(s, 0, 0, 8.12, 7.5, "203D4D", "315A6A")
    nav(s, v[0], color=T["paper"])
    text(s, v[1], .65, 1.27, 6.83, .91, 39, color=T["paper"])
    text(s, v[2], .65, 2.24, 6.83, 1.0, 47, color=T["paper"], bold=True)
    for index in range(3):
        y = 3.88 + index*.86
        text(s, v[3+index*2], .65, y, 1.35, .42, 22, color="EEB9AF")
        text(s, v[4+index*2], 2.25, y, 5.17, .42, 22, color=T["paper"])
    text(s, v[9], 8.72, 3.27, 3.99, .97, 41)
    return s


LAYOUTS = {
    1: cover, 2: route, 3: quiz, 4: principles, 5: approved.scene,
    6: clone_body, 7: timeline, 8: quiz, 9: answer, 10: clone_body,
    11: approved.data, 12: three_actions, 13: two_scenes, 14: clone_body,
    15: quiz, 16: three_scenes, 17: approved.framework, 18: leaflet,
    19: single_scene, 20: single_scene, 21: urgent, 22: kit,
    23: recap, 24: three_scenes, 25: quiz, 26: answer, 27: closing,
}


def pictures(slide):
    return Counter((s.name, hashlib.sha256(s.image.blob).hexdigest())
                   for s in slide.shapes if s.shape_type == 13 and
                   not s.name.startswith("STYLE_"))


def verify(original, result):
    assert len(result.slides) == len(original.slides) == 33
    frame_ids = []
    for page, (old, new) in enumerate(zip(original.slides, result.slides), 1):
        before = Counter(canonical(t) for t in original_texts(old))
        after = Counter(canonical(t) for t in original_texts(new))
        assert before == after, (page, before-after, after-before)
        assert old.notes_slide.notes_text_frame.text == new.notes_slide.notes_text_frame.text, page
        assert pictures(old) == pictures(new), (page, 'image identity')
        assert old._element.get("show", "1") == new._element.get("show", "1"), page
        ids = [s.shape_id for s in new.shapes]
        assert len(ids) == len(set(ids)), (page, 'shape IDs')
        if page <= 27:
            for obj in new.shapes:
                assert obj.left >= 0 and obj.top >= 0, (page, obj.name)
                assert obj.left+obj.width <= result.slide_width+9144, (page, obj.name)
                assert obj.top+obj.height <= result.slide_height+9144, (page, obj.name)
            frame_ids.extend(s.name for s in new.shapes if s.shape_type == 13
                             and not s.name.startswith("STYLE_"))
    assert len(frame_ids) == len(set(frame_ids)) == 18
    chart = result.slides[10]
    for prefix, total, highlighted in (("DATA_GRID_", 100, 28), ("DATA_ED_", 5, 1)):
        group = [s for s in chart.shapes if s.name.startswith(prefix)]
        assert len(group) == total and sum(s.name.endswith("HIGHLIGHT") for s in group) == highlighted


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    WORK.mkdir(parents=True, exist_ok=True)
    originals = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in (ORIGINAL, MANUSCRIPT)}
    original = Presentation(ORIGINAL)
    result = Presentation(ORIGINAL)  # Preserve the original theme and notes conventions.
    old_ids = list(result.slides._sldIdLst)
    for page in range(1, 28):
        LAYOUTS[page](result, original.slides[page-1])
    for old in list(original.slides)[27:]:
        fresh = result.slides.add_slide(result.slide_layouts[6])
        fresh._element.set("show", "0")
        clone_shapes(old, fresh)
        fresh.notes_slide.notes_text_frame.text = old.notes_slide.notes_text_frame.text
    for entry in old_ids:
        result.part.drop_rel(entry.rId)
        result.slides._sldIdLst.remove(entry)
    verify(original, result)
    result.save(OUT/PPT_NAME)
    manuscript = Document(MANUSCRIPT)
    old_paragraphs = [p.text for p in manuscript.paragraphs]
    old_tables = [[[c.text for c in row.cells] for row in t.rows] for t in manuscript.tables]
    manuscript.paragraphs[0].text = old_paragraphs[0] + "\n精致编辑版"
    manuscript.paragraphs[1].text = old_paragraphs[1].replace("合理用药_双人35分钟.pptx", PPT_NAME)
    manuscript.save(OUT/DOC_NAME)
    revised_doc = Document(OUT/DOC_NAME)
    assert [p.text for p in revised_doc.paragraphs][2:] == old_paragraphs[2:]
    assert [[[c.text for c in row.cells] for row in t.rows] for t in revised_doc.tables] == old_tables
    for p, expected in originals.items(): assert hashlib.sha256(p.read_bytes()).hexdigest() == expected
    report = {"main_slides": 27, "total_slides": 33, "hidden_slides": 6,
              "duration_minutes": 35, "originals_unchanged": True,
              "all_page_text_notes_and_picture_pixels_preserved": True,
              "unique_story_and_scientific_frames": 18, "data_grid": [100,28], "data_ed": [5,1],
              "layout_families": {str(p): f.__name__ for p,f in LAYOUTS.items()},
              "pptx": str((OUT/PPT_NAME).relative_to(ROOT)),
              "docx": str((OUT/DOC_NAME).relative_to(ROOT))}
    (WORK/"structural_checks.json").write_text(json.dumps(report,ensure_ascii=False,indent=2))
    print(json.dumps(report,ensure_ascii=False))


if __name__ == "__main__":
    main()
