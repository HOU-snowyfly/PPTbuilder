"""Skill-based cover and late-deck presentation; all text stays editable."""
from pathlib import Path
from pptx.util import Inches
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from visuals import Canvas, INK, CORAL, WHITE, TEAL, LILAC

NAVY='153E75'; PALE='EEF4FA'; ROSE='FFF0F1'; SOFT='F6F8FB'
ASSET_DIR=Path(__file__).parent/'assets'/'anime'
# Measured gutter boundaries of the new sheet; crop natively in PowerPoint.
ACTION_ROWS=[(0,305),(311,617),(624,1024)]
ACTION_COLS=[(0,507),(514,1023),(1031,1536)]
ARTWORKS=[
 ('cover_rational.png','封面：一起核对药盒与说明书','2026-10-13'),
 ('storyboard.png','宿舍、课堂与聚餐分镜','2026-10-12'),
 ('unique_scenes.png','核药时间线、四道题与组胺肖像','2026-10-12'),
 ('actions_storyboard.png','查、核、问、求助、记录与综合案例','2026-10-13'),
]
# Stable content ID, file, source panel ID, meaning; verified against actual PPT pictures.
ART_USAGE=[(1,'cover_rational.png','FULL','封面核药场景')]
ART_USAGE += [(n,'storyboard.png',f'R{r+1}C{c+1}',name) for n,r,c,name in [
 (6,0,0,'头痛'),(6,0,1,'室友递药'),(6,0,2,'核对两盒'),
 (4,1,0,'鼻部舒服'),(4,1,1,'课堂犯困'),
 (18,2,0,'想起已服药'),(18,2,1,'室友劝酒'),(18,2,2,'核实要求')]]
ART_USAGE += [(n,'unique_scenes.png',f'R{r+1}C{c+1}',name) for n,r,c,name in [
 (8,0,0,'查看记录'),(8,0,1,'症状仍在'),(8,0,2,'核对包装'),
 (2,1,0,'开场题'),(9,1,1,'成分题'),(19,1,2,'聚餐题'),
 (28,2,0,'综合题'),(17,2,1,'鼻部组胺肖像')]]
ART_USAGE += [(n,'actions_storyboard.png',f'R{r+1}C{c+1}',name) for n,r,c,name in [
 (20,0,0,'看核问记'),(21,0,1,'读说明书'),(22,0,2,'核使用要求'),
 (23,1,0,'咨询药师'),(24,1,1,'及时求助'),(25,1,2,'留信息和记录'),
 (27,2,0,'已经服药'),(27,2,1,'想加第二盒'),(27,2,2,'准备聚餐')]]

def artwork(c,filename,x,y,w,h,row=None,col=None,focus=.5):
    if row is None:l,t,r,b=0,0,1536,1024;name='ANIME_COVER'
    else:
        l,r=ACTION_COLS[col];t,b=ACTION_ROWS[row]
        name=f'ANIME_ACTIONS_R{row+1}_C{col+1}'
    sw,sh=r-l,b-t;ratio=w/h
    if sw/sh>ratio:
        cut=sw-sh*ratio;l+=cut*focus;r-=cut*(1-focus)
    else:
        cut=sh-sw/ratio;t+=cut*.10;b-=cut*.90
    pic=c.s.shapes.add_picture(str(ASSET_DIR/filename),Inches(x),Inches(y),Inches(w),Inches(h))
    pic.crop_left=l/1536;pic.crop_right=1-r/1536
    pic.crop_top=t/1024;pic.crop_bottom=1-b/1024
    pic.name=name
    pic._element.xpath('.//p:cNvPr')[0].set('descr','原创动漫情境，生成日期与分镜用途见素材附页；对白独立可编辑')
    return pic

def action(c,row,col,x,y,w,h,focus=.5):
    return artwork(c,'actions_storyboard.png',x,y,w,h,row,col,focus)

def speech(c,t,x,y,w,h=1.12,size=23,fill=PALE):
    c.box(x,y,w,h,fill)
    tip=c.shape(MSO_SHAPE.ISOSCELES_TRIANGLE,x+.30,y+h-.02,.25,.18,fill);tip.rotation=180
    c.text(t,x+.22,y+.12,w-.44,h-.24,size,INK,True)

def cover(c,d):
    c.s.background.fill.solid();c.s.background.fill.fore_color.rgb=RGBColor.from_string(WHITE)
    c.box(.65,.46,.08,.28,CORAL)
    c.text('苏州大学 · 校园用药课堂',.87,.40,5.38,.43,17,NAVY,True)
    c.box(6.20,1.23,6.48,4.59,PALE)
    artwork(c,'cover_rational.png',6.32,1.40,6.24,4.16)
    c.text('合理用药',.79,1.69,5.39,1.24,74,NAVY,True)
    c.box(.84,3.04,.82,.08,CORAL)
    c.text('给宿舍药箱上堂课',.83,3.30,5.17,.57,29,NAVY,True)
    c.box(.84,4.23,4.99,1.28,PALE)
    c.text('药名熟悉，就能放心吃？\n先学会看、核、问、记。',1.06,4.39,4.55,.94,23,NAVY,True)
    for j,t in enumerate(['看说明书','核对成分','问清疑点','留好记录']):
        x=.84+j*3.03;c.box(x,6.07,2.84,.63,ROSE if j==1 else PALE)
        c.text(t,x+.18,6.16,2.48,.42,22,NAVY,True,True)
    c.text(d['items'][2],.86,7.00,6.71,.25,15,NAVY)
    c.text('35分钟 · 双人展示',9.42,7.00,3.06,.25,15,NAVY,False,True)

def framework(c,d):
    action(c,0,0,.85,2.14,4.48,2.99)
    speech(c,'拿不准，先别靠感觉。',.85,5.37,4.48,.65,23)
    bodies=['症状、说明书、\n就医信号','成分、用法、禁忌\n与相互作用','带齐信息，\n问医生或药师','药名、时间、\n用量与反应']
    for j,item in enumerate(d['items']):
        x=5.63+(j%2)*3.49;y=2.14+(j//2)*2.06
        c.box(x,y,3.33,1.91,PALE if j%2==0 else ROSE)
        c.text(item.split('｜')[0],x+.26,y+.19,2.81,.58,34,NAVY,True)
        c.text(bodies[j],x+.26,y+.99,2.81,.69,21,INK)

def leaflet(c,d):
    for j,item in enumerate(d['items']):
        x=.85+(j%2)*4.10;y=2.13+(j//2)*1.36
        title,body=item.split('｜');c.box(x,y,3.90,1.21,PALE)
        c.text(title,x+.20,y+.17,3.50,.33,20,NAVY,True)
        c.text(body,x+.20,y+.60,3.50,.43,20,INK)
    action(c,0,1,9.23,2.13,3.22,3.95,focus=.82)

def usage(c,d):
    action(c,0,2,.85,2.13,5.63,3.39)
    speech(c,'用法找到了，别凭感觉改。',.85,5.66,5.63,.54,22)
    for j,item in enumerate(d['items']):
        title,body=item.split('｜');y=2.13+j*1.38
        c.box(6.83,y,5.63,1.22,ROSE if j==2 else PALE)
        c.text(title,7.08,y+.16,5.12,.37,25,NAVY,True)
        c.text(body,7.08,y+.64,5.12,.40,21,INK)

def consult(c,d):
    action(c,1,0,.85,2.13,5.63,3.39)
    speech(c,'药盒和记录，我都带来了。',.85,5.66,5.63,.54,22)
    for j,item in enumerate(d['items']):speech(c,item,6.83,2.13+j*1.38,5.63,1.14,23)

def urgent(c,d):
    c.box(.85,2.13,5.98,1.35,CORAL)
    c.text('120',1.05,2.22,1.87,.85,54,WHITE,True,True)
    c.text('立即急救',1.10,3.09,1.77,.27,18,WHITE,True,True)
    c.text('呼吸困难、喉舌肿胀\n或意识异常',3.15,2.40,3.36,.74,21,WHITE,True)
    for j,(title,body) in enumerate([('疑似超量 / 误服','及时寻求医疗帮助'),('症状持续、加重或反复','及时就诊')]):
        y=3.72+j*1.26;c.box(.85,y,5.98,1.10,PALE)
        c.text(title,1.11,y+.16,5.46,.34,24,NAVY,True)
        c.text(body,1.11,y+.63,5.46,.30,21,INK)
    action(c,1,1,7.15,2.13,5.30,3.39)
    speech(c,'带上药盒与已服记录。',7.15,5.66,5.30,.54,22)

def kit(c,d):
    action(c,1,2,.85,2.13,5.63,3.39)
    speech(c,'需要时，能认清、查到、说清。',.85,5.66,5.63,.54,21)
    for j,item in enumerate(d['items']):
        title,body=item.split('｜');y=2.13+j*1.38
        c.box(6.83,y,5.63,1.22,ROSE if j==2 else PALE)
        c.text(title,7.08,y+.16,5.12,.37,25,NAVY,True)
        c.text(body,7.08,y+.64,5.12,.40,21,INK)

def recap(c,d):
    for j,item in enumerate(d['items']):
        situation,act=item.split(' → ');y=2.13+j*1.02
        c.box(.85,y,11.60,.91,PALE if j%2==0 else ROSE)
        c.text(situation,1.13,y+.22,3.21,.47,23,NAVY,True)
        c.arrow(4.76,y+.32,.45,CORAL)
        c.text(act,5.58,y+.22,6.57,.47,23,NAVY,True)

def final_case(c,d):
    for j,item in enumerate(d['items']):
        x=.82+j*4.03;action(c,2,j,x,2.14,3.79,2.52)
        c.caption(item,x,4.83,3.79,1.12,21)

def final_quiz(c,d):
    for j,item in enumerate(d['items']):
        y=2.17+j*1.34;speech(c,'',.85,y,8.57,1.03,23)
        c.oval(1.08,y+.25,.55,.55,NAVY)
        c.text(item[0],1.08,y+.29,.55,.45,20,WHITE,True,True)
        c.text(item[2:],1.96,y+.16,7.17,.71,23,INK)
    c.box(9.69,2.13,2.76,4.05,PALE)
    c.box(10.04,2.34,2.06,.42,NAVY)
    c.text('可多选',10.12,2.38,1.90,.32,16,WHITE,True,True)
    c.scene(2,0,9.88,2.98,2.38,2.96,extra=True)

def final_answer(c,d):
    c.box(.85,2.13,3.42,4.05,NAVY)
    c.text('A＋B',1.10,2.56,2.92,1.10,54,WHITE,True,True)
    c.text('两步一起做',1.10,3.88,2.92,.48,25,WHITE,True,True)
    # Editable water and information icons support the two selected actions.
    c.box(1.39,4.77,.69,.90,WHITE);c.box(1.45,5.17,.57,.42,'B9DDED')
    c.box(2.54,4.72,1.12,1.01,WHITE)
    for k in range(3):c.line(2.69,4.94+k*.25,3.48,4.94+k*.25,NAVY,2)
    for j,item in enumerate(d['items']):
        title,body=item.split('｜');y=2.13+j*1.38
        c.box(4.61,y,7.84,1.22,ROSE if j==2 else PALE)
        c.text(title,4.89,y+.15,7.28,.38,25,NAVY,True)
        c.text(body,4.89,y+.65,7.28,.38,21,INK)

def closing(c,d):
    c.s.background.fill.solid();c.s.background.fill.fore_color.rgb=RGBColor.from_string(WHITE)
    c.box(.65,.46,.08,.28,CORAL)
    c.text('合理用药',.87,.40,5.38,.43,20,NAVY,True)
    c.text('药可以常备，\n判断不能省略。',.83,1.17,7.09,1.61,43,NAVY,True)
    for j,item in enumerate(d['items']):
        y=3.26+j*1.02;c.box(.85,y,6.78,.83,ROSE if j==1 else PALE)
        tag=['看','核·记','问'][j];c.text(tag,1.01,y+.20,1.19,.43,19,CORAL,True,True)
        c.text(item,2.39,y+.21,4.94,.42,23,NAVY,True)
    # Unique editable editorial illustration, not a reused scene or photograph.
    c.oval(8.07,1.23,4.49,4.49,PALE)
    c.box(9.57,1.71,2.13,2.39,LILAC)
    for k in range(4):c.line(9.86,2.16+k*.36,11.35,2.16+k*.36,NAVY,2)
    c.medbox(8.28,2.20,1.39,2.05,NAVY,'成分')
    c.box(8.65,3.85,3.82,1.87,NAVY)
    c.box(8.84,4.04,3.44,1.49,WHITE)
    c.box(9.45,4.65,1.01,.22,CORAL);c.box(9.84,4.26,.22,1.00,CORAL)
    c.box(9.85,3.49,1.40,.50,NAVY);c.box(10.02,3.64,1.06,.24,WHITE)
    c.text('谢谢大家',8.58,6.18,3.93,.56,27,NAVY,True,True)
    c.text('会看 · 会核 · 会问 · 会记',.88,6.88,7.29,.37,22,NAVY,True)

DRAWERS={1:cover,20:framework,21:leaflet,22:usage,23:consult,24:urgent,25:kit,26:recap,27:final_case,28:final_quiz,29:final_answer,30:closing}
def draw_late_slide(slide,d):
    draw=DRAWERS.get(d['n'])
    if draw is None:return False
    draw(Canvas(slide),d)
    return True
