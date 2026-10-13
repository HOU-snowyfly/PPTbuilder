"""Editable, colorful classroom comics and diagrams for the revised slides."""
import math
from pathlib import Path
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement

INK='203354'; TEAL='008C82'; PURPLE='7952D5'; ORANGE='FFBE45'
CORAL='EB5969'; MINT='E4F8EF'; LILAC='F1EAFF'; CREAM='FFF2D3'
SKIN='FFD3B3'; WHITE='FFFFFF'; GREY='DDE5EF'
FONT='Noto Sans CJK SC'
ANIME_SHEET=Path(__file__).parent/'assets'/'anime'/'storyboard.png'
EXTRA_SHEET=Path(__file__).parent/'assets'/'anime'/'unique_scenes.png'
# Bounds omit the generated sheet gutters; cropping stays native to PowerPoint.
ANIME_ROWS=[(0,304),(313,616),(625,1024)]
ANIME_COLS=[(0,507),(514,1021),(1029,1536)]

class Canvas:
    def __init__(self,slide):self.s=slide
    def shape(self,kind,x,y,w,h,fill=WHITE,stroke=None,lw=1.5):
        o=self.s.shapes.add_shape(kind,Inches(x),Inches(y),Inches(w),Inches(h))
        o.fill.solid();o.fill.fore_color.rgb=RGBColor.from_string(fill)
        o._element.spPr.append(OxmlElement('a:effectLst'))
        for effect in o._element.xpath('.//a:effectRef'):effect.set('idx','0')
        if stroke:o.line.color.rgb=RGBColor.from_string(stroke);o.line.width=Pt(lw)
        else:o.line.fill.background()
        if kind==MSO_SHAPE.ROUNDED_RECTANGLE:o.adjustments[0]=.10
        return o
    def box(self,x,y,w,h,fill=WHITE,stroke=None):return self.shape(MSO_SHAPE.ROUNDED_RECTANGLE,x,y,w,h,fill,stroke)
    def oval(self,x,y,w,h,fill=WHITE,stroke=None):return self.shape(MSO_SHAPE.OVAL,x,y,w,h,fill,stroke)
    def line(self,x1,y1,x2,y2,c=INK,lw=2):
        o=self.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
        o.line.color.rgb=RGBColor.from_string(c);o.line.width=Pt(lw)
        o._element.spPr.append(OxmlElement('a:effectLst'))
        for effect in o._element.xpath('.//a:effectRef'):effect.set('idx','0')
        return o
    def text(self,t,x,y,w,h,size=20,c=INK,bold=False,center=False):
        o=self.s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=o.text_frame
        tf.word_wrap=True;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=Pt(0)
        tf.vertical_anchor=MSO_ANCHOR.MIDDLE
        for i,row in enumerate(t.split('\n')):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=row;p.line_spacing=1.05;p.space_after=Pt(0)
            p.alignment=PP_ALIGN.CENTER if center else PP_ALIGN.LEFT
            for r in p.runs:r.font.name=FONT;r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=RGBColor.from_string(c)
        return o
    def arrow(self,x,y,w=.45,c=CORAL):return self.shape(MSO_SHAPE.CHEVRON,x,y,w,.30,c)
    def badge(self,t,x,y,w,c=TEAL):self.box(x,y,w,.36,c);self.text(t,x+.07,y+.02,w-.14,.30,12,WHITE,True,True)
    def bubble(self,t,x,y,w,h=1.0,size=19):
        self.box(x,y,w,h,WHITE)
        tip=self.shape(MSO_SHAPE.ISOSCELES_TRIANGLE,x+.25,y+h-.01,.28,.24,WHITE);tip.rotation=180
        self.text(t,x+.16,y+.08,w-.32,h-.16,size,INK,True)
    def sparkle(self,x,y,w=.22,c=ORANGE):self.shape(MSO_SHAPE.STAR_4_POINT,x,y,w,w,c)
    def pill(self,x,y,w=1,h=.34,c=CORAL):
        self.box(x,y,w,h,c);self.box(x+.04,y+.04,w*.44,h-.08,WHITE)
    def medbox(self,x,y,w=1,h=1.2,c=TEAL,label='药盒'):
        self.box(x,y,w,h,WHITE,c);self.box(x+.05,y+.05,w-.10,.22,c)
        self.line(x+w*.5,y+.40,x+w*.5,y+.70,c,3);self.line(x+w*.35,y+.55,x+w*.65,y+.55,c,3)
        self.text(label,x+.05,y+h-.34,w-.10,.25,12,c,True,True)
    def scene(self,row,col,x,y,w,h,window=None,extra=False):
        left,right=ANIME_COLS[col];top,bottom=ANIME_ROWS[row]
        if window:
            a,b,z,q=window;sw=right-left;sh=bottom-top
            left,right=left+a*sw,left+z*sw;top,bottom=top+b*sh,top+q*sh
        sw,sh=right-left,bottom-top;ratio=w/h
        if sw/sh>ratio:
            cut=(sw-sh*ratio)/2;left+=cut;right-=cut
        else:
            cut=sh-sw/ratio;top+=cut*.10;bottom-=cut*.90
        pic=self.s.shapes.add_picture(str(EXTRA_SHEET if extra else ANIME_SHEET),Inches(x),Inches(y),Inches(w),Inches(h))
        pic.crop_left=left/1536;pic.crop_right=1-right/1536
        pic.crop_top=top/1024;pic.crop_bottom=1-bottom/1024
        pic.name=f"ANIME_{'EXTRA' if extra else 'BASE'}_R{row+1}_C{col+1}"
        pic._element.xpath('.//p:cNvPr')[0].set('descr','原创AI生成动漫情境；人物和事件为虚构教学示意')
        return pic
    def caption(self,t,x,y,w,h=1.00,size=20):
        self.box(x,y,w,h,'F1F5FA')
        self.text(t,x+.18,y+.09,w-.36,h-.18,size,INK,True)
    def student(self,x,y,k=.8,mood='happy',roommate=False):
        return self.scene(2,1,x,y,1.65*k,2.10*k,extra=True)

    def books(self,x,y,w=1.2):
        self.box(x,y,w,.20,TEAL);self.box(x+.08,y-.24,w,.20,PURPLE);self.box(x+.15,y-.48,w,.20,CORAL)
        for a,b in [(x+.09,y+.07),(x+.17,y-.17),(x+.24,y-.41)]:self.line(a,b,a+w-.28,b,WHITE,1.5)
    def brain(self,x,y,k=1):
        # Soft lobes and short curved-looking folds; icon, not anatomical scale.
        for a,b,w,h in [(0,.30,.65,.72),(.12,.05,.75,.76),(.58,0,.76,.80),(1.08,.22,.67,.77),(.60,.52,.77,.57)]:
            self.oval(x+a*k,y+b*k,w*k,h*k,'D8C4FF',PURPLE)
        for a,b,z,q in [(.31,.28,.48,.47),(.70,.24,.66,.57),(1.14,.38,1.30,.59),(.43,.76,.69,.85),(1.04,.80,1.18,.93)]:
            self.line(x+a*k,y+b*k,x+z*k,y+q*k,PURPLE,2)

def comic(c,d):
    for j,item in enumerate(d['items']):
        x=.82+j*4.03
        c.scene(0,j,x,2.14,3.79,2.52)
        c.caption(item.split('｜',1)[-1],x,4.83,3.79,1.12,21)

def ingredient(c):
    for x,fill,name,sub in [(1.02,'F0F4FA','药盒 A','单方退热止痛药'),(7.13,'F0F4FA','药盒 B','复方感冒药')]:
        c.box(x,2.12,5.20,3.19,fill);c.text(name,x+.27,2.31,2.15,.42,25,INK,True);c.text(sub,x+.27,2.80,4.6,.31,17)
        c.box(x+.24,3.31,4.72,.58,CORAL);c.text('对乙酰氨基酚',x+.34,3.38,4.52,.42,24,WHITE,True,True)
    c.text('只有这一种有效成分',1.30,4.34,4.6,.39,19,INK,True,True)
    for text,y,color in [('马来酸氯苯那敏',4.04,INK),('盐酸伪麻黄碱',4.61,INK)]:
        c.box(7.37,y,4.72,.45,color);c.text(text,7.48,y+.04,4.50,.35,19,WHITE,True,True)
    c.line(5.98,3.60,6.43,3.60,CORAL,2.5);c.line(6.43,3.60,6.43,5.22,CORAL,2.5)
    c.line(7.37,3.60,6.90,3.60,CORAL,2.5);c.line(6.90,3.60,6.90,5.22,CORAL,2.5)
    c.line(6.43,5.22,6.67,5.40,CORAL,2.5);c.line(6.90,5.22,6.67,5.40,CORAL,2.5)
    c.box(4.30,5.40,4.75,.44,'FFEEF0');c.text('同一成分，来自两盒',4.40,5.43,4.55,.36,22,CORAL,True,True)
    c.text('① 圈出重复　→　② 核规格与已服情况　→　③ 未确认，不自行叠加',1.1,5.97,11.1,.30,17,INK,True,True)

def timeline(c,d):
    for j,item in enumerate(d['items']):
        x=.82+j*4.03;c.box(x,2.12,3.79,3.96,'F1F5FA')
        if j<2:c.badge('18:00' if j==0 else '20:00',x+.22,2.29,1.35,INK)
        c.scene(0,j,x+.20,2.84,3.39,2.12,extra=True)
        c.text(item.split('｜',1)[-1],x+.22,5.15,3.35,.72,19,INK,True)

def microbes(c):
    c.box(.85,2.08,5.63,4.02,MINT);c.box(6.83,2.08,5.63,4.02,LILAC)
    c.badge('细菌',1.12,2.31,1.03,TEAL);c.badge('病毒',7.10,2.31,1.03,PURPLE)
    # Bacterium: membrane, genetic material and dots representing ribosomes.
    c.oval(1.34,3.21,2.26,1.36,TEAL);c.oval(1.45,3.32,2.04,1.14,'B5EFDF')
    for j in range(12):
        a=j*math.tau/12;cx=2.47+1.14*math.cos(a);cy=3.89+.70*math.sin(a)
        c.line(cx,cy,cx+.22*math.cos(a),cy+.22*math.sin(a),TEAL,2)
    for j in range(5):c.line(1.85+j*.22,3.85+(j%2)*.15,2.07+j*.22,3.85+((j+1)%2)*.15,PURPLE,3)
    for a,b in [(1.72,3.58),(2.25,3.54),(2.97,3.66),(1.78,4.17),(2.61,4.23)]:c.oval(a,b,.11,.11,ORANGE)
    c.text('有细胞结构',3.81,3.48,2.27,.39,23,TEAL,True);c.text('在适宜条件下\n可自行繁殖',3.81,3.99,2.27,.61,17)
    c.text('某些细菌结构或过程\n可成为抗生素的靶点',1.15,5.11,4.99,.67,21,INK,True)
    # Virus: shell and genome, then a host cell rather than a free-standing cell.
    cx,cy=8.04,3.90
    for j in range(10):
        a=j*math.tau/10;x=cx+.65*math.cos(a);y=cy+.65*math.sin(a)
        c.line(cx+.43*math.cos(a),cy+.43*math.sin(a),x,y,PURPLE,2);c.oval(x-.06,y-.06,.12,.12,ORANGE)
    c.oval(cx-.48,cy-.48,.96,.96,PURPLE);c.oval(cx-.34,cy-.34,.68,.68,'D8C4FF')
    for j in range(3):c.line(cx-.18+j*.13,cy-.16,cx-.05+j*.13,cy+.15,CORAL,2)
    c.arrow(9.02,3.79,.40,PURPLE);c.oval(9.65,3.24,2.15,1.40,'DBF4ED',TEAL);c.oval(10.28,3.66,.61,.54,TEAL)
    c.text('宿主细胞',9.77,4.77,1.96,.30,17,TEAL,True,True)
    c.text('无细胞结构\n依赖宿主细胞复制',7.15,5.12,4.99,.66,21,INK,True)

def allergy_scene(c):
    for j,x in enumerate([.85,6.83]):
        c.scene(1,j,x,2.12,5.63,2.98)
        c.caption(['终于不用一直擦鼻涕！','鼻子好了，脑子想“下课”……'][j],x,5.25,5.63,.60,22)

def quiz(c,d):
    # All options share size, weight and spacing; no visual answer cue.
    for j,item in enumerate(d['items']):
        y=2.12+j*1.25;c.box(.84,y,9.03,1.03,'F4F6FA')
        c.badge(item[0],1.08,y+.33,.59,INK);c.text(item[2:],1.96,y+.17,7.57,.70,23,INK,False)
    c.box(10.14,2.12,2.32,3.55,'F4F6FA')
    row,col={2:(1,0),9:(1,1),19:(1,2),28:(2,0)}[d['n']]
    c.scene(row,col,10.40,2.97,1.82,2.31,extra=True)
    c.badge('可多选' if d.get('question_mode')=='多选' else '单选',10.69,2.35,1.21,PURPLE)


def principles(c,d):
    for j,item in enumerate(d['items']):
        x=.85+(j%2)*6.00;y=2.12+(j//2)*2.06
        key,body=item.split('｜',1)
        c.box(x,y,5.62,1.77,'EEF4FA' if j%2==0 else 'FFF0F1')
        c.badge(f'0{j+1}',x+.25,y+.22,.56,INK)
        c.text(key,x+1.01,y+.17,4.22,.48,29,INK,True)
        c.text(body,x+.29,y+1.00,4.99,.48,23)

def antibiotic_risk(c):
    for x in [.85,6.83]:c.box(x,2.10,5.63,4.10,'F0F4FA')
    c.text('美国门诊抗生素处方',1.13,2.31,5.00,.44,24,INK,True)
    c.text('至少',1.13,3.38,2.37,.36,22,CORAL,True)
    c.text('28%',1.13,3.89,2.37,.86,48,CORAL,True)
    for j in range(100):
        x=3.80+(j%10)*.225;y=3.21+(j//10)*.225
        c.box(x,y,.18,.18,CORAL if j<28 else 'D6E0ED')
    c.text('被估计为不必要',1.13,5.55,5.00,.41,24,CORAL,True)
    c.text('药物副作用导致的急诊',7.11,2.31,5.00,.44,24,INK,True)
    c.text('5次中，约1次',7.11,3.03,5.00,.53,32,CORAL,True)
    c.text('涉及抗生素反应',7.11,3.67,5.00,.38,23,INK,True)
    for j in range(5):
        x=7.14+j*.96;y=4.40;color=CORAL if j==0 else INK
        c.box(x,y,.76,1.01,color)
        c.box(x+.24,y+.15,.28,.08,WHITE);c.box(x+.34,y+.05,.08,.28,WHITE)
        c.box(x+.30,y+.68,.16,.33,WHITE)
        for k in range(2):c.box(x+.12+k*.36,y+.45,.12,.12,WHITE)
    c.text('皮疹、腹泻，也可能严重过敏',7.11,5.66,5.00,.33,19,INK,True)

def antibiotic_actions(c,d):
    for j,item in enumerate(d['items']):
        x=.85+j*4.04;c.box(x,2.13,3.80,3.93,'EEF4FA')
        title,body=item.split('｜',1);c.badge(f'0{j+1}',x+.26,2.34,.60,INK)
        if j==0:
            c.medbox(x+.57,3.03,1.00,1.25,INK,'抗生素')
            c.oval(x+1.85,3.10,1.03,1.03,'FFEEF0',CORAL)
            c.text('需要？',x+1.95,3.45,.82,.29,17,CORAL,True,True)
            c.line(x+2.66,3.93,x+3.12,4.34,CORAL,6)
        elif j==1:
            c.medbox(x+.47,3.07,1.01,1.29,INK,'室友剩药')
            c.medbox(x+2.19,3.07,1.01,1.29,INK,'我的用药')
            c.oval(x+1.50,3.43,.60,.60,'FFEEF0',CORAL)
            c.line(x+1.61,3.55,x+1.99,3.93,CORAL,3)
        else:
            c.box(x+.67,3.02,2.30,1.54,WHITE,INK)
            for k,t in enumerate(['剂量','间隔','疗程']):
                yy=3.17+k*.43;c.text(t,x+1.28,yy,1.36,.28,16)
                c.line(x+.83,yy+.13,x+.96,yy+.24,CORAL,2)
                c.line(x+.96,yy+.24,x+1.16,yy+.01,CORAL,2)
        c.text(title,x+.26,4.86,3.28,.39,24,INK,True)
        c.text(body,x+.26,5.36,3.28,.56,17)

def drinking_case(c,d):
    for j,item in enumerate(d['items']):
        x=.82+j*4.03;c.scene(2,j,x,2.14,3.79,2.52)
        c.caption(item.split('｜',1)[-1],x,4.83,3.79,1.12,21)

def agenda(c,d):
    for j,item in enumerate(d['items']):
        x=.85+j*3.03;c.box(x,2.16,2.83,3.93,'EEF4FA' if j%2==0 else 'FFF0F1')
        c.badge(f'0{j+1}',x+.24,2.38,.63,INK)
        if j==0:
            c.oval(x+.59,3.21,1.17,1.17,WHITE,INK)
            c.line(x+1.57,4.18,x+2.11,4.69,CORAL,5)
            c.text('?',x+.82,3.40,.71,.70,35,CORAL,True,True)
        elif j==1:c.medbox(x+.75,3.15,1.33,1.53,INK,'宿舍药箱')
        elif j==2:
            c.box(x+.59,3.13,1.65,1.59,WHITE,INK)
            for k in range(3):
                y=3.41+k*.40;c.line(x+.77,y,x+.91,y+.13,CORAL,2);c.line(x+.91,y+.13,x+1.13,y-.08,CORAL,2);c.line(x+1.33,y,x+2.04,y,INK,2)
        else:
            for k,t in enumerate(['A','B','C']):c.badge(t,x+.43+k*.66,3.66,.56,CORAL if k==1 else INK)
        title,time=item.split('｜');c.text(title[3:],x+.20,5.07,2.43,.43,25,INK,True,True)
        c.text(time,x+.20,5.63,2.43,.31,18,INK,False,True)

def dose_rules(c,d):
    for j,item in enumerate(d['items']):
        x=.85+j*4.04;c.box(x,2.13,3.80,3.93,'EEF4FA' if j!=2 else 'FFF0F1')
        title,body=item.split('｜');c.badge(f'0{j+1}',x+.26,2.35,.60,INK)
        if j==0:
            c.oval(x+.78,3.03,1.53,1.53,WHITE,INK)
            c.line(x+1.55,3.31,x+1.55,3.79,INK,3);c.line(x+1.55,3.79,x+1.99,4.02,CORAL,3)
            c.text('?',x+2.45,3.28,.58,.93,40,CORAL,True,True)
        elif j==1:
            c.box(x+.78,3.03,2.07,1.56,WHITE,INK)
            for k,t in enumerate(['剂量','间隔','疗程']):c.text(t,x+1.02,3.19+k*.41,1.56,.29,18,INK,True)
        else:
            c.oval(x+.75,3.09,2.24,1.40,'FFEEF0',CORAL)
            c.oval(x+.99,3.31,1.76,.96,WHITE,INK)
            c.line(x+1.87,3.47,x+1.87,4.10,INK,2)
            c.text('完整保护层',x+.71,4.62,2.38,.31,17,CORAL,True,True)
        c.text(title,x+.26,5.10,3.28,.39,25,INK,True)
        c.text(body,x+.26,5.59,3.28,.51,18)

def finalcase(c,d):
    c.scene(2,2,.85,2.13,5.63,3.97,extra=True)
    for j,item in enumerate(d['items']):
        y=2.13+j*1.38;c.box(6.83,y,5.63,1.22,'EEF4FA' if j!=2 else 'FFF0F1')
        c.badge(f'0{j+1}',7.07,y+.38,.60,INK)
        c.text(item,7.91,y+.16,4.27,.90,22,INK,True)

def draw_revised_slide(slide,d):
    from medical_visuals import histamine_medical
    if histamine_medical(slide,d):return True
    from late_visuals import draw_late_slide
    if draw_late_slide(slide,d):return True
    c=Canvas(slide);n=d['n']
    if d['layout']=='quiz':quiz(c,d)
    elif n==4:allergy_scene(c)
    elif n==3:principles(c,d)
    elif n==6:comic(c,d)
    elif n==7:ingredient(c)
    elif n==8:timeline(c,d)
    elif n==11:microbes(c)
    elif n==13:antibiotic_risk(c)
    elif n==14:antibiotic_actions(c,d)
    elif n==5:agenda(c,d)
    elif n==22:dose_rules(c,d)
    elif n==27:finalcase(c,d)
    elif n==18:drinking_case(c,d)
    else:return False
    return True
