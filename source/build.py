from pathlib import Path
import math, json, csv, re
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_MARK
from pptx.oxml.xmlchemy import OxmlElement
from docx import Document
from docx.shared import Inches as DI, Pt as DP, RGBColor as DC
from docx.oxml import OxmlElement as DX
from docx.oxml.ns import qn
from content import SLIDES, SOURCES

ROOT=Path(__file__).parent
OUT=ROOT/'交付文件'; OUT.mkdir(exist_ok=True)
PHOTODIR=ROOT/'assets'/'photos'
B='153E75'; R='E3464F'; W='FFFFFF'; FONT='Noto Sans CJK SC'
PHOTO_CREDITS=[
 ('tablet-hand.jpg','Towfiqu barbhuiya','https://www.pexels.com/photo/hand-holding-a-pill-15703464/','Pexels License','https://www.pexels.com/license/'),
 ('pill-bottle.jpg','Pexels photo 16051935','https://www.pexels.com/photo/16051935/','Pexels License','https://www.pexels.com/license/'),
 ('college-campus.jpg','George Pak','https://www.pexels.com/photo/7972512/','Pexels License','https://www.pexels.com/license/'),
 ('evening.jpg','LZ Jian','https://www.pexels.com/photo/black-lamp-beside-bed-5968800/','Pexels License','https://www.pexels.com/license/'),
 ('pharmacy-01.jpg','Árpád Czapp','https://unsplash.com/photos/a-room-filled-with-lots-of-shelves-filled-with-boxes-and-boxes-tvP6pCnq9iI','Unsplash License','https://unsplash.com/license'),
 ('pharmacy-02.jpg','National Cancer Institute','https://unsplash.com/photos/pharmacist-reaching-for-medication-on-shelf-byGTytEGjBo','Unsplash License','https://unsplash.com/license'),
 ('pharmacy-03.jpg','paws and prints','https://unsplash.com/photos/sQjHIOhiL0E','Unsplash License','https://unsplash.com/license'),
 ('hospital-02.jpg','Tasha Kostyuk','https://unsplash.com/photos/empty-hospital-corridor-with-benches-and-doors-Pk-KuizxQv8','Unsplash License','https://unsplash.com/license'),
 ('clinic-01.jpg','Martha Dominguez de Gouveia','https://unsplash.com/photos/hospital-lobby-reception-with-signage-nMyM7fxpokE','Unsplash License','https://unsplash.com/license'),
]
PHOTO_BY_PAGE={1:['tablet-hand.jpg'],2:['tablet-hand.jpg'],5:['college-campus.jpg','pharmacy-01.jpg','pharmacy-02.jpg','hospital-02.jpg'],6:['pill-bottle.jpg'],9:['pill-bottle.jpg'],14:['pharmacy-03.jpg','pharmacy-02.jpg','pill-bottle.jpg'],15:['college-campus.jpg'],16:['evening.jpg'],18:['pill-bottle.jpg'],19:['pharmacy-03.jpg'],21:['pill-bottle.jpg'],22:['pill-bottle.jpg','pharmacy-01.jpg','tablet-hand.jpg'],23:['pharmacy-02.jpg'],24:['hospital-02.jpg'],25:['pharmacy-01.jpg'],27:['tablet-hand.jpg'],28:['pharmacy-01.jpg'],30:['clinic-01.jpg']}
PHOTO_INDEX={item[0]:item for item in PHOTO_CREDITS}
P=Presentation(); P.slide_width=Inches(13.333333); P.slide_height=Inches(7.5)
P.core_properties.title='合理用药：给宿舍药箱上堂课'
P.core_properties.subject='苏州大学选修课｜双人35分钟课堂展示'
P.core_properties.author='课堂展示制作'
PARTS={1:'01 / 基本原则',2:'02 / 宿舍误区',3:'03 / 安全用药',4:'04 / 选择挑战'}

def color(c): return RGBColor.from_string(c)
def shape(s,typ,x,y,w,h,fill=W,line=B,lw=1.5):
    o=s.shapes.add_shape(typ, Inches(x), Inches(y), Inches(w), Inches(h))
    o.fill.solid(); o.fill.fore_color.rgb=color(fill)
    o.line.color.rgb=color(line); o.line.width=Pt(lw)
    o._element.spPr.append(OxmlElement('a:effectLst'))
    for ef in o._element.xpath('.//a:effectRef'): ef.set('idx','0')
    return o
def rect(s,x,y,w,h,fill=W,line=B,lw=1.5,round=False):
    return shape(s,MSO_SHAPE.ROUNDED_RECTANGLE if round else MSO_SHAPE.RECTANGLE,x,y,w,h,fill,line,lw)
def circle(s,x,y,d,fill=W,line=B,lw=1.5): return shape(s,MSO_SHAPE.OVAL,x,y,d,d,fill,line,lw)
def line(s,x1,y1,x2,y2,c=B,lw=2):
    o=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2));o.line.color.rgb=color(c);o.line.width=Pt(lw);o._element.spPr.append(OxmlElement('a:effectLst'))
    for ef in o._element.xpath('.//a:effectRef'):ef.set('idx','0')
    return o
def txt(s,t,x,y,w,h,size=24,c=B,bold=False,align=PP_ALIGN.LEFT,val=MSO_ANCHOR.MIDDLE):
    o=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=o.text_frame
    tf.clear();tf.word_wrap=True;tf.margin_left=Pt(0);tf.margin_right=Pt(0);tf.margin_top=Pt(0);tf.margin_bottom=Pt(0);tf.vertical_anchor=val
    for idx,part in enumerate(t.split('\n')):
        p=tf.paragraphs[0] if idx==0 else tf.add_paragraph();p.text=part;p.alignment=align;p.space_before=Pt(0);p.space_after=Pt(4);p.line_spacing=1.08
        for r in p.runs:
            r.font.name=FONT;r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=color(c)
            rp=r._r.get_or_add_rPr();ea=OxmlElement('a:ea');ea.set('typeface',FONT);rp.append(ea)
    return o
def pill(s,x,y,w=1.3,h=.48,rot=-25):
    o=rect(s,x,y,w,h,fill=R,line=R,round=True);o.rotation=rot
    # A white inset makes this an illustrated capsule, all colors remain within palette.
    z=circle(s,x+.1,y+.09,h*.6,W,W);z.rotation=rot
    return o
def cross(s,x,y,size=.5,c=R):
    rect(s,x+size*.33,y,size*.34,size,c,c,0);rect(s,x,y+size*.33,size,size*.34,c,c,0)
def medbox(s,x,y,w=1.1,h=1.7,label='药',accent=R):
    rect(s,x,y,w,h,W,B,2,True);rect(s,x,y,w,min(.28,h*.2),accent,accent,0)
    cross(s,x+w*.32,y+h*.32,min(w*.35,h*.24),accent);txt(s,label,x+.06,y+h*.74,w-.12,h*.19,min(13,max(9,h*9)),B,True,PP_ALIGN.CENTER)
def person(s,x,y,scale=1,sleepy=False):
    circle(s,x+.42*scale,y,.96*scale,W,B,3)
    # hair and a swept fringe
    shape(s,MSO_SHAPE.ARC,x+.39*scale,y-.04*scale,1.02*scale,.76*scale,B,B,2)
    for ex in [.7,1.08]:
        if sleepy: line(s,x+ex*scale,y+.48*scale,x+(ex+.14)*scale,y+.48*scale,B,2)
        else: circle(s,x+ex*scale,y+.43*scale,.075*scale,B,B,0)
    line(s,x+.83*scale,y+.71*scale,x+1.01*scale,y+.71*scale,R,2)
    rect(s,x+.15*scale,y+1.03*scale,1.5*scale,1.5*scale,B,B,1,True)
    line(s,x+.69*scale,y+1.06*scale,x+.9*scale,y+1.31*scale,W,2)
    line(s,x+1.11*scale,y+1.06*scale,x+.9*scale,y+1.31*scale,W,2)
    rect(s,x+.55*scale,y+1.81*scale,.67*scale,.31*scale,B,W,1,True)
    if sleepy:txt(s,'Z z',x+1.48*scale,y-.02*scale,.9*scale,.4*scale,22,R,True)
def cup(s,x,y,sz=1,wine=False):
    if wine:
        shape(s,MSO_SHAPE.TRAPEZOID,x,y,.8*sz,.7*sz,W,R,2)
        line(s,x+.4*sz,y+.7*sz,x+.4*sz,y+1.2*sz,R,2);line(s,x+.15*sz,y+1.2*sz,x+.65*sz,y+1.2*sz,R,2)
    else:
        rect(s,x,y,.8*sz,1.0*sz,W,B,2,True);line(s,x+.1*sz,y+.38*sz,x+.7*sz,y+.38*sz,B,2)
        shape(s,MSO_SHAPE.ARC,x+.61*sz,y+.2*sz,.5*sz,.5*sz,W,B,2)
def book(s,x,y,w=1.6):
    rect(s,x,y,w,.28,R,R,0,True);rect(s,x+.05,y+.08,w-.12,.12,W,W,0)
    rect(s,x+.12,y-.32,w,.28,B,B,0,True);rect(s,x+.17,y-.24,w-.12,.12,W,W,0)
def clock(s,x,y,d=.8):
    circle(s,x,y,d,W,B,2);line(s,x+d/2,y+d/2,x+d/2,y+d*.2,R,2);line(s,x+d/2,y+d/2,x+d*.73,y+d*.61,B,2)
def arrow(s,x,y,w=.55,c=R): return shape(s,MSO_SHAPE.CHEVRON,x,y,w,.42,c,c,0)
def badge(s,t,x,y,w=1.5,c=R):rect(s,x,y,w,.36,c,c,0,True);txt(s,t,x+.08,y+.02,w-.16,.28,12,W,True,PP_ALIGN.CENTER)
def photo(s,name,x,y,w,h,credit=True):
    """Place a real photograph with an undistorted, center-cropped view."""
    path=PHOTODIR/name
    with Image.open(path) as im: iw,ih=im.size
    pic=s.shapes.add_picture(str(path),Inches(x),Inches(y),Inches(w),Inches(h))
    source_ratio=iw/ih; box_ratio=w/h
    if source_ratio>box_ratio:
        cut=(1-box_ratio/source_ratio)/2;pic.crop_left=cut;pic.crop_right=cut
    else:
        cut=(1-source_ratio/box_ratio)/2;pic.crop_top=cut;pic.crop_bottom=cut
    frame=rect(s,x,y,w,h,W,B,1.5);frame.fill.background()
    if credit:
        rect(s,x+.02,y+h-.35,w-.04,.33,W,W,0)
        txt(s,'真实照片 · 场景示意',x+.14,y+h-.32,w-.28,.26,10,B,True)
    return pic
def base(d,dark=False):
    s=P.slides.add_slide(P.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=color(B if dark else W)
    if not dark:
        rect(s,.55,.4,.07,.27,R,R,0);txt(s,PARTS[d['part']],.75,.32,4,.45,13,B,True)
        txt(s,d['title'],.65,.98,12.05,.78,30,B,True)
        line(s,.65,6.39,12.68,6.39,B,.7)
        txt(s,d['takeaway'],.7,6.5,11.75,.47,17,R,True)
        txt(s,('来源：'+d['refs']+'  ·  完整链接见附页') if d['refs'] else '合理用药｜课堂科普',.7,7.12,10.5,.2,9,B)
        txt(s,f"{d['n']:02d} / 30",11.75,7.06,.9,.28,11,B,False,PP_ALIGN.RIGHT)
    # speaker notes include full source URLs, not just numbers
    ss='\n'.join(f'[{sid}] {title}\n{url}' for sid,title,url in SOURCES if sid in d['refs'].split('; '))
    photo_notes='\n'.join(f"{PHOTO_INDEX[n][1]}｜{PHOTO_INDEX[n][2]}｜{PHOTO_INDEX[n][3]}" for n in PHOTO_BY_PAGE.get(d['n'],[]))
    s.notes_slide.notes_text_frame.text=f"第{d['n']}页｜{d['speaker']}主讲｜建议{d['sec']}秒\n{d.get('cue','')}\n\n{d['notes']}\n\n视觉说明：{d['visual']}\n\n参考：\n{ss}\n\n真实照片为图库场景示意，不呈现虚构案例中的真实人物、药品或事件。\n图片来源：\n{photo_notes}"
    return s
def panel(s,x,y,w,h,title,body,num=None):
    rect(s,x,y,w,h,W,B,1.4,True)
    if num:badge(s,num,x+.22,y+.22,.58)
    txt(s,title,x+.28,y+.73 if num else y+.32,w-.56,.58,25,B,True)
    txt(s,body,x+.28,y+1.47 if num else y+1.09,w-.56,h-(1.65 if num else 1.2),20,B)
def textrows(s,items,x=.9,y=2.1,w=11.6,step=1.14,size=24):
    for j,t in enumerate(items):
        circle(s,x,y+j*step+.15,.14,R,R,0);txt(s,t,x+.36,y+j*step,w-.36,.87,size,B)
def splititem(t):
    if '｜' in t:return t.split('｜',1)
    return t,''

for d in SLIDES:
    kind=d['layout'];s=base(d,dark=kind in ['cover','closing'])
    it=d['items']
    if kind=='cover':
        badge(s,'校园用药生存课',.7,.65,2.2,R)
        txt(s,'合理用药',.72,1.55,7,1.18,61,W,True)
        txt(s,'给宿舍药箱上堂课',.78,2.86,6.7,.76,32,W,True)
        txt(s,it[0],.8,4.02,6.7,.53,23,W)
        txt(s,it[1]+'\n'+it[2],.8,5.5,6.9,.97,17,W)
        photo(s,'tablet-hand.jpg',8.1,1.16,4.55,5.35)
        txt(s,'35分钟 · 2位讲者 · 4部分',.8,7.03,10,.28,12,W)
    elif kind=='closing':
        badge(s,'把判断带回生活',.8,.72,2.3)
        txt(s,'药可以常备，\n判断不能省略。',.8,1.65,7.35,1.9,43,W,True)
        for j,t in enumerate(it):txt(s,t,.9,4.05+j*.55,7.2,.42,21,W)
        txt(s,'会看 · 会核 · 会问 · 会记',.9,6.18,7.5,.52,25,W,True)
        photo(s,'clinic-01.jpg',8.65,1.55,3.8,4.88)
        txt(s,'谢谢大家',10.65,6.68,2,.36,17,W)
    elif kind=='quiz':
        quiz_photo={2:'tablet-hand.jpg',9:'pill-bottle.jpg',19:'pharmacy-03.jpg',28:'pharmacy-01.jpg'}[d['n']]
        for j,t in enumerate(it):
            rect(s,.8,2.02+j*1.27,9.24,1.02,W,B,1.6,True)
            circle(s,1.03,2.21+j*1.27,.62,B,B,0)
            txt(s,t[0],1.03,2.24+j*1.27,.62,.52,23,W,True,PP_ALIGN.CENTER)
            txt(s,t[2:],1.95,2.13+j*1.27,7.79,.79,23,B)
        photo(s,quiz_photo,10.27,2.02,2.25,3.56)
    elif kind=='four':
        for j,t in enumerate(it):
            a,b=splititem(t);x=.8+(j%2)*6.04;y=1.98+(j//2)*2.05
            rect(s,x,y,5.73,1.8,W,B,1.4,True)
            txt(s,a,x+.25,y+.2,4.9,.57,30,R,True)
            txt(s,b,x+.25,y+.94,5.18,.52,22,B)
    elif kind=='agenda':
        for j,t in enumerate(it):
            x=.78+j*3.17
            rect(s,x,2.24,2.96,3.45,W,B,1.4,True)
            photo(s,['college-campus.jpg','pharmacy-01.jpg','pharmacy-02.jpg','hospital-02.jpg'][j],x+.12,2.36,2.72,1.32,False)
            a,b=t.split('｜');txt(s,a[:2],x+.25,3.78,2.5,.55,34,R,True)
            txt(s,a[3:],x+.25,4.45,2.5,.5,25,B,True);txt(s,b,x+.25,5.06,2.5,.38,18,B)
    elif kind=='case':
        photo(s,{6:'pill-bottle.jpg',15:'college-campus.jpg',16:'evening.jpg'}[d['n']],.77,2.03,4.06,3.93)
        for j,t in enumerate(it):
            a,b=splititem(t)
            if not b and '：' in t:a,b=t.split('：',1)
            x=5.23;y=2.04+j*1.3
            txt(s,a,x,y,7.1,.47,22,R,True);txt(s,b if b else t,x,y+.49,7.1,.66,24,B)
    elif kind=='labels':
        for j in range(2):
            x=.85+j*6.25
            rect(s,x,2.1,5.55,2.56,W,B,2,True)
            badge(s,'教学标签 '+('A' if j==0 else 'B'),x+.23,2.32,1.7,B)
            txt(s,'有效成分',x+.26,2.98,4.9,.4,18,B)
            rect(s,x+.22,3.55,5.1,.68,W,R,2,True)
            txt(s,'对乙酰氨基酚',x+.4,3.64,4.7,.46,27,R,True)
            if j==1:txt(s,'＋其他成分',x+.29,4.3,4.8,.27,14,B)
        arrow(s,3.7,5.01,.5);arrow(s,9.45,5.01,.5)
        txt(s,'成分会汇合，风险不能按药盒“分开算”',1.1,5.45,11,.6,25,B,True,PP_ALIGN.CENTER)
    elif kind=='errorchart':
        txt(s,'美国成人模拟任务｜n = 500｜2012年｜单位：%',.88,1.93,11.4,.45,18,B)
        data=CategoryChartData();data.categories=['单一产品模拟超量','模拟联用重复超量'];data.add_series('占参与者比例',[23.8,45.6])
        chart=s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED,Inches(.9),Inches(2.62),Inches(11.5),Inches(2.75),data).chart
        chart.has_legend=False;chart.has_title=False
        chart.font.name=FONT;chart.font.size=Pt(17);chart.font.color.rgb=color(B)
        va=chart.value_axis;va.minimum_scale=0;va.maximum_scale=100;va.major_unit=20;va.tick_labels.number_format='0"%"';va.tick_labels.font.name=FONT;va.tick_labels.font.size=Pt(14);va.tick_labels.font.color.rgb=color(B);va.has_major_gridlines=False
        va.format.line.color.rgb=color(B);va.major_tick_mark=XL_TICK_MARK.NONE;va.minor_tick_mark=XL_TICK_MARK.NONE
        ca=chart.category_axis;ca.tick_labels.font.name=FONT;ca.tick_labels.font.size=Pt(18);ca.tick_labels.font.color.rgb=color(B);ca.format.line.fill.background();ca.major_tick_mark=XL_TICK_MARK.NONE
        plot=chart.plots[0];plot.gap_width=70;plot.has_data_labels=True;dl=plot.data_labels;dl.position=XL_LABEL_POSITION.OUTSIDE_END;dl.number_format='0.0';dl.font.name=FONT;dl.font.size=Pt(23);dl.font.bold=True;dl.font.color.rgb=color(R)
        for j,point in enumerate(chart.series[0].points):point.format.fill.solid();point.format.fill.fore_color.rgb=color(B if j==0 else R);point.format.line.fill.background()
        txt(s,'研究判断的是潜在超量风险，不是实际中毒事件。',.95,5.67,11.6,.38,19,B)
    elif kind=='amrchart':
        txt(s,'2019年全球细菌耐药负担｜单位：万人｜2022年发表',.85,1.92,11.6,.45,18,B)
        x0=3.2;span=8.7
        for j,(name,val,lo,hi,c) in enumerate([('相关死亡',495,362,657,B),('归因死亡',127,91.1,171,R)]):
            yy=2.83+j*1.34
            txt(s,name,.86,yy-.03,2.13,.49,23,B,True)
            rect(s,x0,yy,span*val/700,.5,c,c,0)
            line(s,x0+span*lo/700,yy+.7,x0+span*hi/700,yy+.7,c,2)
            for e in [lo,hi]:line(s,x0+span*e/700,yy+.6,x0+span*e/700,yy+.8,c,2)
            txt(s,str(val),x0+span*val/700+.14,yy-.03,1.2,.48,23,c,True)
            txt(s,f'区间：{lo:g}—{hi:g}',.87,yy+.56,2.24,.55,13,B)
        line(s,x0,5.62,x0+span,5.62,B,1)
        for tick in range(0,701,100):
            x=x0+span*tick/700;line(s,x,5.62,x,5.7,B,1);txt(s,str(tick),x-.28,5.79,.56,.25,12,B,False,PP_ALIGN.CENTER)
        txt(s,'细线表示95%不确定性区间',.9,6.06,5,.25,13,B)
    elif kind=='microbes':
        for j in range(2):rect(s,.83+j*6.16,2.04,5.78,3.9,W,B,1.4,True)
        shape(s,MSO_SHAPE.ROUNDED_RECTANGLE,1.43,2.78,2.03,1.07,W,B,3)
        for a in range(5):circle(s,1.68+a*.29,3.08,.12,R,R,0)
        for a in range(5):line(s,1.5+a*.41,2.71,1.47+a*.41,2.45,B,2)
        txt(s,'细菌',3.78,2.87,2,.56,30,B,True)
        txt(s,'具有细胞结构\n细胞壁合成等可成为靶点',1.18,4.29,5.15,1.04,22,B)
        cx,cy=8.35,3.15
        for a in range(12):
            th=a*math.pi/6;line(s,cx+.51*math.cos(th),cy+.51*math.sin(th),cx+.81*math.cos(th),cy+.81*math.sin(th),R,2)
            circle(s,cx+.81*math.cos(th)-.05,cy+.81*math.sin(th)-.05,.1,R,R,0)
        circle(s,cx-.51,cy-.51,1.02,W,R,3);pill(s,cx-.24,cy-.1,.48,.19)
        txt(s,'病毒',9.66,2.87,2,.56,30,B,True)
        txt(s,'没有细胞结构\n抗生素不治疗普通感冒病毒',7.31,4.29,5.08,1.04,22,B)
    elif kind=='selection':
        groups=[[(j%4,j//4,R if j in [3,10] else B) for j in range(12)],[(0,0,R),(2,1,R)],[(j%4,j//4,R) for j in range(12)]]
        for j,g in enumerate(groups):
            x=.85+j*4.21;rect(s,x,2.1,3.55,3.72,W,B,1.4,True)
            txt(s,['原有差异','药物选择压力','存活并繁殖'][j],x+.2,2.32,3.15,.53,23,B,True)
            for xx,yy,cc in g:
                o=rect(s,x+.38+xx*.7,3.2+yy*.57,.43,.25,cc,cc,0,True);o.rotation=20*(xx%2)
            if j==1:txt(s,'敏感菌更易被抑制',x+.19,5.1,3.18,.42,17,B)
            if j<2:arrow(s,x+3.68,3.72,.35)
        txt(s,'蓝：敏感菌　红：耐药菌　｜　数量仅为示意，不代表实验比例',.93,5.94,11.5,.27,14,B)
    elif kind=='target':
        if d['n']==4:
            pill(s,1.68,2.88,1.4,.6,0)
            txt(s,'药物',1.17,4.18,2.6,.5,24,B,True,PP_ALIGN.CENTER);arrow(s,3.55,3.1,.66)
            shape(s,MSO_SHAPE.HEXAGON,5.1,2.69,1.44,1.22,W,B,3)
            txt(s,'受体 / 酶',4.5,4.18,2.8,.5,24,B,True,PP_ALIGN.CENTER);arrow(s,7.5,3.1,.66)
            circle(s,9.19,2.77,1.12,R,R,0);txt(s,'信号',9.23,3.05,1.04,.4,20,W,True,PP_ALIGN.CENTER)
            txt(s,'反应发生变化',8.4,4.18,3.15,.5,24,B,True,PP_ALIGN.CENTER)
            txt(s,'蛋白质的结构，与功能有关。',1.1,5.28,11.1,.61,27,B,True,PP_ALIGN.CENTER)
        else:
            for j,(a,b) in enumerate([('外周组织','组胺信号 → 参与过敏症状'),('脑内','组胺信号 → 参与清醒状态')]):
                x=.85+j*6.17;rect(s,x,2.11,5.73,3.7,W,B,1.5,True);badge(s,a,x+.24,2.39,1.53,B)
                circle(s,x+.51,3.34,.65,R,R,0);arrow(s,x+1.57,3.45,.53);shape(s,MSO_SHAPE.HEXAGON,x+2.69,3.26,.91,.85,W,B,2)
                txt(s,b,x+.26,4.51,5.19,.84,23,B)
            txt(s,'部分抗组胺药进入脑内后，可能引起嗜睡。',1,5.97,11.5,.28,15,R,True)
    elif kind=='three':
        for j,t in enumerate(it):
            a,b=splititem(t);x=.82+j*4.2
            rect(s,x,2.18,3.77,3.66,W,B,1.4,True)
            photo(s,(['pharmacy-03.jpg','pharmacy-02.jpg','pill-bottle.jpg'] if d['n']==14 else ['pill-bottle.jpg','pharmacy-01.jpg','tablet-hand.jpg'])[j],x+.13,2.3,3.51,1.35,False)
            badge(s,f'0{j+1}',x+.23,3.8,.58)
            txt(s,a,x+.23,4.24,3.3,.46,22,B,True)
            txt(s,b,x+.23,4.86,3.28,.7,18,B)
    elif kind=='answer':
        txt(s,'C' if d['n']==10 else 'B',.82,2.25,2.3,2.3,120,R,True,PP_ALIGN.CENTER)
        textrows(s,it,x=3.45,y=2.12,w=9.03,step=1.23,size=25)
    elif kind=='drinks':
        photo(s,'pill-bottle.jpg',1.04,2.17,4.65,3.58)
        for j,t in enumerate(it):
            a,b=splititem(t);txt(s,a,6.06,2.08+j*1.32,6.34,.49,22,B,True);txt(s,b,6.06,2.65+j*1.32,6.34,.53,20,B)
    elif kind=='leaflet':
        rect(s,1.02,1.97,11.27,4.12,W,B,2,True)
        for j,t in enumerate(it):
            x=1.29+(j%2)*4.03;y=2.23+(j//2)*1.19
            circle(s,x,y,.55,R,R,0);txt(s,str(j+1),x,y+.03,.55,.45,20,W,True,PP_ALIGN.CENTER)
            txt(s,t,x+.72,y-.01,3.2,.67,19,B,True)
        photo(s,'pill-bottle.jpg',9.42,2.16,2.55,3.49)
        txt(s,'请以手中的真实说明书为准',1.25,5.83,7.8,.2,12,B)
    elif kind=='chat':
        photo(s,'pharmacy-02.jpg',.9,2.16,2.75,3.72)
        rect(s,3.94,2.16,8.45,3.72,W,B,1.5,True)
        textrows(s,it,x=4.25,y=2.3,w=7.79,step=1.08,size=23)
    elif kind=='urgent':
        rect(s,.85,2.12,3.12,3.79,R,R,0,True);txt(s,'120',1.1,2.8,2.6,1.08,64,W,True,PP_ALIGN.CENTER);txt(s,'严重急症\n立即求助',1.18,4.36,2.43,1.08,25,W,True,PP_ALIGN.CENTER)
        photo(s,'hospital-02.jpg',10.45,2.14,2.0,3.74)
        for j,t in enumerate(it):
            a,b=splititem(t);txt(s,a,4.39,2.17+j*1.25,5.85,.51,20,B,True);txt(s,b,4.39,2.75+j*1.25,5.85,.42,19,R)
    elif kind=='kit':
        photo(s,'pharmacy-01.jpg',1,2.24,3.6,3.3)
        textrows(s,it,x=5.03,y=2.29,w=7.43,step=1.22,size=24)
    elif kind=='recap':
        for j,t in enumerate(it):
            a,b=t.split(' → ');yy=2.0+j*.81
            txt(s,f'{j+1:02d}',.84,yy,.69,.49,23,R,True)
            txt(s,a,1.76,yy,4.46,.54,23,B,True);arrow(s,6.46,yy+.06,.4);txt(s,b,7.31,yy,5.01,.54,23,B)
    elif kind=='finalcase':
        photo(s,'tablet-hand.jpg',.87,2.08,2.75,3.67)
        for j,t in enumerate(it):
            a,b=splititem(t);yy=2.08+j*1.28;rect(s,4.0,yy,8.46,1.04,W,B,1.5,True);txt(s,t,4.28,yy+.12,7.89,.78,24,B)

# Reference slides are hidden during the timed presentation, but remain editable and exportable.
for page,start in enumerate(range(0,len(SOURCES),5),1):
    s=P.slides.add_slide(P.slide_layouts[6]);s._element.set('show','0')
    s.background.fill.solid();s.background.fill.fore_color.rgb=color(W)
    txt(s,f'参考资料 / {page}',.7,.52,11,.65,31,B,True)
    txt(s,'附页，不计入35分钟｜点击条目标题可访问原文｜检索日期：2026-10-09',.72,1.26,11.9,.41,15,R)
    for j,(sid,title,url) in enumerate(SOURCES[start:start+5]):
        yy=1.96+j*.96;badge(s,sid,.77,yy,.7,B)
        o=txt(s,title,1.73,yy-.06,10.88,.72,14.5,B)
        for p in o.text_frame.paragraphs:
            for run in p.runs:run.hyperlink.address=url
        line(s,1.73,yy+.76,12.59,yy+.76,B,.45)
    txt(s,'图表按原始数据重绘；案例为虚构，实拍照片另见图片来源附页。',.78,7.02,11.9,.3,12,B)
    s.notes_slide.notes_text_frame.text='参考文献附页，不逐页讲解。\n\n'+'\n\n'.join(f'[{a}] {b}\n{c}' for a,b,c in SOURCES[start:start+5])

for page,start in enumerate(range(0,len(PHOTO_CREDITS),5),1):
    s=P.slides.add_slide(P.slide_layouts[6]);s._element.set('show','0')
    s.background.fill.solid();s.background.fill.fore_color.rgb=color(W)
    txt(s,f'真实照片来源 / {page}',.7,.52,11,.65,31,B,True)
    txt(s,'图片为真实图库照片，仅作课堂场景示意；照片中的人物、药物、场所与虚构案例无关。',.72,1.25,11.9,.6,15,R)
    for j,(name,creator,url,license_name,license_url) in enumerate(PHOTO_CREDITS[start:start+5]):
        yy=1.96+j*.95
        badge(s,f'P{start+j+1}',.77,yy,.7,B)
        o=txt(s,f'{creator} · {name}',1.73,yy-.06,10.88,.42,14.5,B)
        for p in o.text_frame.paragraphs:
            for run in p.runs:run.hyperlink.address=url
        o2=txt(s,license_name+' · 点击标题访问原图',1.73,yy+.37,10.88,.32,11,R)
        for p in o2.text_frame.paragraphs:
            for run in p.runs:run.hyperlink.address=license_url
        line(s,1.73,yy+.79,12.59,yy+.79,B,.45)
    txt(s,'附页，不计入35分钟｜完整原图链接与许可见配套文件。',.78,7.02,11.9,.3,12,B)
    s.notes_slide.notes_text_frame.text='真实照片来源附页，不逐页讲解。\n\n'+'\n\n'.join(f'{name}｜{creator}\n{url}\n{license_name}: {license_url}' for name,creator,url,license_name,license_url in PHOTO_CREDITS[start:start+5])

P.save(OUT/'合理用药_双人35分钟.pptx')

# Create speaker manuscript and a separate production / source guide.
def setupdoc(title):
    doc=Document();sec=doc.sections[0];sec.top_margin=DI(.7);sec.bottom_margin=DI(.65);sec.left_margin=DI(.78);sec.right_margin=DI(.78)
    normal=doc.styles['Normal'];normal.font.name=FONT;normal.font.size=DP(10.5);normal.font.color.rgb=DC.from_string(B)
    normal._element.rPr.rFonts.set(qn('w:eastAsia'),FONT);normal.paragraph_format.space_after=DP(6);normal.paragraph_format.line_spacing=1.25
    for name in ['Title','Heading 1','Heading 2','Heading 3']:
        st=doc.styles[name];st.font.name=FONT;st.font.color.rgb=DC.from_string(B);st._element.rPr.rFonts.set(qn('w:eastAsia'),FONT)
    doc.add_heading(title,0)
    f=sec.footer.paragraphs[0];f.text='合理用药｜苏州大学选修课课堂展示　 ·　';field=DX('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
    return doc
def mm(v):return f'{v//60:02d}:{v%60:02d}'
def table(doc,heads,rows):
    t=doc.add_table(rows=1, cols=len(heads));t.style='Light Shading Accent 1'
    for c,h in zip(t.rows[0].cells,heads):c.text=h
    for row in rows:
        for c,v in zip(t.add_row().cells,row):c.text=str(v)
    return t
def sourceparas(doc):
    for sid,title,url in SOURCES:
        doc.add_paragraph(f'[{sid}] {title}');p=doc.add_paragraph(url);p.paragraph_format.space_after=DP(10)

doc=setupdoc('合理用药\n双人逐页演讲稿')
doc.add_paragraph('配套文件：合理用药_双人35分钟.pptx｜30页主讲＋3页文献附页＋2页图片来源附页')
doc.add_paragraph('目标：面向具有高中生物基础的非医学专业大一学生，学会处理常见用药选择并识别误区。全部小苏故事均为虚构情境；真实研究单独标明。')
doc.add_heading('35分钟时间与分工',1)
table(doc,['部分','主讲页','时间','内容'],[['第一部分','1—5','00:00—05:00','基本原则'],['第二部分','6—19','05:00—22:00','五类宿舍误区'],['第三部分','20—26','22:00—30:00','安全用药方法'],['第四部分','27—30','30:00—35:00','简单综合选择题与收尾']])
doc.add_paragraph('甲主讲第1—16页，计划16分55秒；乙主讲第17—30页，计划18分05秒。第29页甲补充约15秒，第30页共同收尾，实际发言时间更接近。第16页末交接，不中途介绍新身份。')
doc.add_paragraph('互动时，一人主持，另一人观察举手并翻页。不要求同学解释观点，不询问私人疾病史。参考资料页设为隐藏，课堂主线不会自动进入附页。')
doc.add_heading('排练与控时',1)
doc.add_paragraph('计时是演讲目标，不是自动播放。稿件约5700个汉字，另有停顿、读题、举手与同桌讨论，按自然讲解语速排练后校准。首次完整计时，优先把总时长控制在34—36分钟；不要为凑时间机械拖慢每句话。')
doc.add_paragraph('到第5页末应约5分钟，第19页末约22分钟，第26页末约30分钟。超时1分钟：压缩第18页重复解释约30秒，并将第19页同桌讨论与点评合计缩短30秒。提前1分钟：按稿中停顿完整读图，补充解释第8页“模拟任务”和第13页“两种口径”，不要临时增加医学结论。')
doc.add_paragraph('使用PowerPoint演示者视图查看备注。翻页采用手动点击，不设自动计时；答案页在投票后再显示。把封面的“同学甲/同学乙”替换成姓名。正式展示前，在实际电脑上检查字体和投影。')
elapsed=0
for d in SLIDES:
    doc.add_heading(f"{d['n']:02d}｜{d['title']}",1)
    doc.add_paragraph(f"主讲：{d['speaker']}　建议：{d['sec']}秒　累计：{mm(elapsed)}—{mm(elapsed+d['sec'])}")
    if d['cue']:doc.add_paragraph('舞台提示：'+d['cue'])
    for para in d['notes'].split('\n'):doc.add_paragraph(para)
    if d['refs']:doc.add_paragraph('本页依据：'+d['refs'])
    elapsed+=d['sec']
doc.add_heading('互动答案速查',1)
table(doc,['题目','题目页','揭晓位置','答案'],[['开场选择','2','3','B'],['重复成分','9','10','C'],['聚餐饮酒','19','19页口头＋20页提示','B'],['综合挑战','28','29','B']])
doc.add_heading('参考资料',1);sourceparas(doc)
doc.save(OUT/'合理用药_双人逐页讲稿.docx')

guide=setupdoc('合理用药\n大纲、逐页设计与资料核验')
guide.add_paragraph('设计依据：PPT制作30分、语言表达30分、内容完整性15分、互动效果15分、时间把握10分。围绕“课堂可读、案例可讲、选择可做、时间可控”落实。')
guide.add_heading('整体方案',1)
table(guide,['项目','实现方式'],[['内容结构','基本原则5分钟；宿舍误区17分钟；安全用药8分钟；选择挑战5分钟。'],['视觉','16:9；模板蓝 #153E75、红 #E3464F、白 #FFFFFF；真实照片保留自然色。标题约30pt，主体以20—30pt为主。'],['图片','18页主讲页加入9张不同的真实图库照片，按版面裁切后嵌入PPT；案例照片仅作场景示意。机制示意仍为可编辑形状。'],['图表','第8页为PowerPoint原生图表；第13页为可编辑矢量图，含不确定性区间。'],['互动','三次穿插投票＋一题综合挑战；A/B/C举手，不要求开放表达。'],['演讲支持','完整逐页备注、双人台词、交接提示和累计时间；附独立Word讲稿。'],['高中生物','蛋白质结构与功能、受体与信号、细菌和病毒、变异与自然选择。'],['文件使用','PPT内含3页隐藏文献附页、2页隐藏图片来源附页；PDF用于预览，编辑请使用PPTX。']])
guide.add_heading('四部分大纲',1)
table(guide,['部分','页码','要点'],[['1 基本原则','1—5','开场情境；合理用药四原则；靶点示意；四段路线。'],['2 宿舍误区','6—19','重复成分；两篇研究图表；抗生素与耐药；备考产品；抗过敏嗜睡；酒精相互作用。'],['3 安全用药','20—26','看核问记；说明书六处；剂量、时间和剂型；咨询信息；就医信号；药箱管理；动作回顾。'],['4 选择挑战','27—30','复方感冒药＋聚餐情境；三选一；逐项揭晓；三句收尾。']])
guide.add_heading('逐页文案与图片 / 图表说明',1)
for d in SLIDES:
    guide.add_heading(f"{d['n']:02d}｜{d['title']}",2)
    guide.add_paragraph('页面短文案：'+'；'.join(d['items']))
    guide.add_paragraph('底部关键句：'+d['takeaway'])
    guide.add_paragraph('视觉：'+d['visual']+('；另嵌入真实照片：'+'、'.join(PHOTO_BY_PAGE[d['n']]) if d['n'] in PHOTO_BY_PAGE else ''))
    guide.add_paragraph(f"分工与时间：{d['speaker']}，{d['sec']}秒。依据：{d['refs'] or '原创课堂组织 / 综合归纳'}。")
guide.add_heading('数据口径与重绘规则',1)
guide.add_paragraph('图1（第8页）：Wolf等，2012年，500名美国成人门诊参与者，模拟非处方对乙酰氨基酚用药任务。单一产品任务中23.8%安排出24小时超过4克的用量；联用任务中45.6%出现重复成分造成的潜在超量。4克为该研究的判定口径，不是本课给出的用药建议，也不能代替中国具体药品说明书的限量。两指标可能重叠，不相加；不能解释为实际中毒、肝损伤或中国大学生发生率。原始数值保留一位小数，条形轴从0至100%。')
guide.add_paragraph('图2（第13页）：Antimicrobial Resistance Collaborators，Lancet 2022；数据年2019，全球细菌耐药负担模型估计。相关死亡495万人（95%不确定性区间362—657万人）；归因死亡127万人（91.1—171万人）。相关口径以无耐药感染为反事实，归因口径以药物敏感感染替代耐药感染为反事实。不是两个人群可以相加，也不表示这些死亡均由个人滥用抗生素造成。没有把2019年数据写成当前最新数据。')
guide.add_paragraph('两组数据适合比较大小，因此采用条形图，而不是不适合当前数据结构的瀑布图。图表均按原文数据重绘并标注来源；没有直接截取论文整图。原始绘图数据另附CSV，便于更新。')
guide.add_heading('教学边界与常见追问',1)
for t in [
 '问：感冒就一定不能用抗生素？答：普通感冒通常由病毒引起，抗生素不能治疗病毒；是否合并细菌感染或属于其他疾病，需要医生评估。',
 '问：抗生素是不是一定要吃到药盒空？答：按处方确定的方案使用；有问题及时联系医生，不自行缩短或延长，不按包装剩余量决定疗程。',
 '问：药盒名字不同为什么不能一起吃？答：商品名不是成分表，同一种有效成分可能出现在多个产品中；具体联用应核对并咨询。',
 '问：非嗜睡抗过敏药是不是绝对不会困？答：嗜睡风险较低不等于人人都没有，仍需看说明书并观察个体反应。',
 '问：能不能给个统一剂量或喝酒等待时长？答：具体药物、规格、病情和个体情况不同，不能给出统一值；看实际说明书和专业建议。',
 '问：耐药是不是身体对药适应了？答：此处讨论的是细菌对抗菌药物的抵抗能力，不是人的身体产生所谓耐药体质。',
 '课堂材料用于一般健康科普，不承担个体诊断和处方功能。案例均为虚构，无真实同学病史；急症行动以中国场景的120表述。']:
    guide.add_paragraph(t)
guide.add_heading('评分点与排练检查',1)
table(guide,['评分项','检查动作'],[['PPT制作 30','后排能读到主体文字；三色统一；图形无遮挡；不直接念参考文献。'],['语言表达 30','用自己的语气练熟讲稿；每页只围绕一个重点；甲乙交接自然；数据解释口径。'],['内容完整性 15','四部分完整；五类案例齐全；原则、机制和行动互相对应。'],['互动效果 15','先给阅读时间，再逐项举手；投票后揭晓；即使无人举手也继续解释。'],['时间把握 10','至少完整计时排练一次；检查5、22、30分钟节点；主线约35分钟。']])
guide.add_heading('完整参考资料',1);sourceparas(guide)
guide.add_heading('真实照片来源与许可',1)
guide.add_paragraph('以下均为真实图库照片，非生成图。图库照片仅说明场景；人物、药品、建筑与小苏虚构案例无关，也不代表苏州大学校园。Pexels 与 Unsplash 许可允许将照片用于演示文稿；保留摄影者与原图链接便于复核。')
for name,creator,url,license_name,license_url in PHOTO_CREDITS:
    guide.add_paragraph(f'{name}｜{creator}｜{license_name}')
    guide.add_paragraph(url+'\n'+license_url)
guide.add_paragraph('检索日期：2026-10-09。网页更新时间可能改变；论文数值按上述固定发表版本使用。全文由课堂语言转述，未照搬整段原文或未授权品牌图像。')
guide.save(OUT/'合理用药_大纲与逐页设计说明.docx')

with (OUT/'图表原始数据与口径.csv').open('w',encoding='utf-8-sig',newline='') as f:
    cw=csv.writer(f);cw.writerow(['页码','指标','数值','单位','下限','上限','口径','来源'])
    cw.writerow([8,'单一产品模拟超量',23.8,'%', '', '', '美国成人n=500；潜在超量；非实际中毒率；可与下一指标重叠','S3'])
    cw.writerow([8,'模拟联用重复超量',45.6,'%', '', '', '美国成人n=500；潜在超量；非实际中毒率；不可相加','S3'])
    cw.writerow([13,'与细菌耐药相关死亡',495,'万人',362,657,'2019全球估计；95%不确定性区间；与归因口径不可相加','S5'])
    cw.writerow([13,'归因于细菌耐药的死亡',127,'万人',91.1,171,'2019全球估计；95%不确定性区间','S5'])

with (OUT/'来源链接与使用说明.txt').open('w',encoding='utf-8') as f:
    f.write('合理用药：给宿舍药箱上堂课\n30页主讲＋3页隐藏文献附页＋2页隐藏图片来源附页，35分钟含互动。\n\n使用步骤：\n1. 用PowerPoint打开PPTX，把封面的同学甲/乙替换为姓名。\n2. 打开演示者视图，备注含逐页讲稿。\n3. 主讲文字、机制示意和图表可编辑；照片可在PPT中替换或裁切。第8页图表右键“编辑数据”；第13页数值改动时需同步调整图形长度和区间。\n4. 全稿使用Noto Sans CJK SC（思源黑体系列）。如果电脑未安装，可在PowerPoint“替换字体”改为微软雅黑；替换后检查换行。PDF用于固定版式预览。\n5. 来源附页设为隐藏，放映时跳过；编辑视图仍可查看。\n6. 至少完整计时排练一次，检查5、22、30分钟节点；35分钟是目标时长，文件不自动播放。\n\n研究数据按原文重绘。校园故事及教学标签为原创示意；真实照片为图库场景示意，不对应虚构故事中的真实人物、药品或场所。\n\n')
    for a,b,c in SOURCES:f.write(f'[{a}] {b}\n{c}\n\n')
    f.write('真实照片来源与许可：\n')
    for name,creator,url,license_name,license_url in PHOTO_CREDITS:
        f.write(f'{name}｜{creator}\n原图：{url}\n许可：{license_name} {license_url}\n\n')
    f.write('检索日期：2026-10-09\n')

with (OUT/'真实照片来源.csv').open('w',encoding='utf-8-sig',newline='') as f:
    writer=csv.writer(f);writer.writerow(['文件','摄影者/来源','原图页面','使用许可','许可链接','使用页码'])
    for name,creator,url,license_name,license_url in PHOTO_CREDITS:
        pages='、'.join(str(n) for n,photos in PHOTO_BY_PAGE.items() if name in photos)
        writer.writerow([name,creator,url,license_name,license_url,pages])

(ROOT/'slides.json').write_text(json.dumps(SLIDES,ensure_ascii=False,indent=2),encoding='utf-8')
print('Created files:',*[str(p) for p in OUT.iterdir()],sep='\n')
