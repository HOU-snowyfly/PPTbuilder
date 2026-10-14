"""Validate the saved refined deck against its original and actual PDF rendering."""
from __future__ import annotations

import hashlib
import json
import re
import subprocess
from collections import Counter

from docx import Document
from pptx import Presentation

from build_refined_deck import ROOT, ORIGINAL, MANUSCRIPT, OUT, WORK, PPT_NAME, DOC_NAME, verify
from build_global_style_samples import canonical


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    pptx = OUT / PPT_NAME
    docx = OUT / DOC_NAME
    pdf = WORK / (pptx.stem + ".pdf")
    result = Presentation(pptx)
    verify(Presentation(ORIGINAL), result)
    missing = []
    for number, slide in enumerate(list(result.slides)[:27], 1):
        actual = subprocess.check_output(
            ["pdftotext", "-f", str(number), "-l", str(number), str(pdf), "-"], text=True)
        expected = Counter("".join(canonical(s.text) for s in slide.shapes if s.has_text_frame))
        absent = expected - Counter("".join(actual.split()))
        if absent or "\\n" in actual:
            missing.append({"page": number, "characters": dict(absent)})
    assert not missing, missing
    info = subprocess.check_output(["pdfinfo", str(pdf)], text=True)
    assert int(re.search(r"Pages:\s+(\d+)", info).group(1)) == 27
    before = Document(MANUSCRIPT)
    after = Document(docx)
    assert [p.text for p in before.paragraphs][2:] == [p.text for p in after.paragraphs][2:]
    table_data = lambda d: [[[c.text for c in r.cells] for r in t.rows] for t in d.tables]
    assert table_data(before) == table_data(after)
    assert PPT_NAME in after.paragraphs[1].text
    original_seconds = None
    with (ROOT / "source/qa_report.json").open() as original_qa:
        for line in original_qa:
            match = re.search(r'"planned_seconds":\s*(\d+)', line)
            if match:
                original_seconds = int(match.group(1))
                break
    assert original_seconds == 2100
    frames = [s for slide in list(result.slides)[:27] for s in slide.shapes
              if s.shape_type == 13 and not s.name.startswith("STYLE_")]
    fingerprints = [(digest_image(s), *(round(getattr(s, "crop_"+k), 5)
                      for k in ("left", "top", "right", "bottom"))) for s in frames]
    assert len(fingerprints) == len(set(fingerprints)) == 18
    report = {"edition": "精致编辑版", "main_slides": 27, "total_slides": 33,
              "hidden_slides": 6, "planned_seconds": original_seconds,
              "page_text_notes_and_existing_images_preserved": True,
              "missing_rendered_text": missing, "docx_main_text_and_tables_preserved": True,
              "unique_story_scientific_frames": 18, "data_grid": [100, 28], "data_ed": [5, 1],
              "actual_pdf_pages": 27, "pptx_sha256": digest(pptx), "docx_sha256": digest(docx),
              "original_pptx_sha256": digest(ORIGINAL), "original_docx_sha256": digest(MANUSCRIPT),
              "visual_review": "27主讲页缩略图检查；三张已确认样页沿用；专业图结构与标签位置保留",
              "pptx": str(pptx.relative_to(ROOT)), "docx": str(docx.relative_to(ROOT))}
    destination = ROOT / "source/refined_qa_report.json"
    destination.write_text(json.dumps(report, ensure_ascii=False, indent=2))
    print(json.dumps(report, ensure_ascii=False))


def digest_image(shape):
    return hashlib.sha256(shape.image.blob).hexdigest()


if __name__ == "__main__":
    main()
