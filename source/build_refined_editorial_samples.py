"""Refine the approved material direction, preserving all original sample content.

Creates three editable 16:9 sample pages and six side-by-side review pages:
original versus refinement, followed by previous material version versus refinement.
"""
from __future__ import annotations

import hashlib
import json
import subprocess
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

from build_editorial_texture_samples import (
    ORIGINAL, PAPER, PAGES, ROOT, THEME, clone_shapes, compare, footer,
    gradient, nav, verify,
)
from build_global_style_samples import box, emergency_icon, original_texts, picture, rgb, text


OUT = ROOT / "work/refined-editorial-samples"
PREVIOUS = ROOT / "work/editorial-texture-samples/合理用药_质感改版_三页样页.pptx"
THEME.update({"paper": "F8F5EE", "tint": "F3EEE4", "muted": "DBD6CA",
              "ink": "234453", "accent": "D65162", "rule": "D5CEC2"})


def base(presentation, original):
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(THEME["paper"])
    backdrop = slide.shapes.add_picture(str(PAPER), 0, 0,
                                       width=presentation.slide_width,
                                       height=presentation.slide_height)
    backdrop.name = "STYLE_WARM_PAPER"
    # Editable color wash quiets the paper without changing its source pixels.
    wash = box(slide, 0, 0, presentation.slide_width / 914400,
               presentation.slide_height / 914400, "FBF8F2")
    wash.name = "STYLE_PAPER_WASH"
    color = wash._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    alpha = OxmlElement("a:alpha")
    alpha.set("val", "64000")
    color.append(alpha)
    slide.notes_slide.notes_text_frame.text = original.notes_slide.notes_text_frame.text
    return slide, original_texts(original)


def scene(presentation, original):
    slide, values = base(presentation, original)
    nav(slide, values[0])
    heading = text(slide, values[1].replace("，", "，\n", 1),
                   .64, .69, 12.02, 1.18, 36, bold=True)
    introduction = heading.text_frame.paragraphs[0]
    introduction.font.size = Pt(24)
    introduction.font.bold = False
    introduction.font.color.rgb = rgb("50646B")
    panels = [s for s in original.shapes if s.shape_type == 13]
    picture(slide, panels[0], 0, 2.13, 6.74, 4.43)
    text(slide, values[4], .49, 6.76, 5.94, .4, 21)
    # Each supporting action and its dialogue now form a compact, aligned unit.
    for panel, caption, x in ((panels[1], values[5], 7.13),
                              (panels[2], values[6], 10.03)):
        box(slide, x, 2.45, 2.63, 2.91, "FCFAF5")
        picture(slide, panel, x, 2.45, 2.63, 1.75)
        text(slide, caption.replace("，", "，\n", 1), x + .13,
             4.43, 2.37, .77, 20)
    gradient(slide, 7.13, 6.25, 5.53, .82, "234655", "355968", angle=0)
    takeaway = text(slide, values[2].replace("，", "，\n", 1), 7.37, 6.34,
                    5.05, .63, 21, color=THEME["paper"])
    takeaway.text_frame.paragraphs[1].font.bold = True
    footer(slide, values[3])
    return slide


def data(presentation, original):
    slide, values = base(presentation, original)
    gradient(slide, 0, 0, 4.1, 7.5, "1D3B4B", "31596A")
    nav(slide, values[0], color=THEME["paper"])
    title = values[1].replace("不需要的抗生素，", "不需要的\n抗生素，\n")
    title = title.replace("真实伤害", "\n真实伤害")
    heading = text(slide, title, .65, 1.1, 2.84, 3.13, 29,
                   color=THEME["paper"])
    for index in (1, 3):
        heading.text_frame.paragraphs[index].font.bold = True
    box(slide, .65, 4.91, .35, .028, THEME["accent"])
    text(slide, values[2].replace("，", "，\n", 1), .65, 5.34,
         2.84, 1.22, 21, color=THEME["paper"])
    metric = values[4:]
    # Read the measured object before the large number and proportion grid.
    text(slide, metric[0], 4.72, .71, 7.94, .43, 24)
    text(slide, metric[1], 4.72, 1.76, 3.9, .32, 20,
         color=THEME["accent"])
    text(slide, metric[2], 4.66, 2.18, 4.23, 1.22, 74,
         color=THEME["accent"], bold=True)
    text(slide, metric[3], 4.72, 3.71, 4.59, .4, 22)
    for index in range(100):
        row, col = divmod(index, 10)
        obj = box(slide, 9.61 + col * .273, 1.13 + row * .273, .222, .222,
                  THEME["accent"] if index < 28 else THEME["muted"], rounded=True)
        obj.name = f"DATA_GRID_{index+1:03d}_{'HIGHLIGHT' if index < 28 else 'REST'}"
    box(slide, 4.72, 4.35, 7.94, .012, THEME["rule"])
    text(slide, metric[4], 4.72, 4.59, 7.94, .4, 22)
    text(slide, metric[5], 4.72, 5.12, 4.75, .59, 31,
         color=THEME["accent"], bold=True)
    text(slide, metric[6], 9.65, 5.19, 3.01, .38, 20)
    for index in range(5):
        emergency_icon(slide, 4.77 + index * 1.6, 5.86, 1.17,
                       THEME["accent"] if index == 0 else THEME["ink"],
                       f"DATA_ED_{index+1}_{'HIGHLIGHT' if index == 0 else 'REST'}")
    text(slide, metric[7], 4.72, 7.01, 6.1, .31, 19)
    footer(slide, values[3])
    return slide


def framework(presentation, original):
    slide, values = base(presentation, original)
    gradient(slide, 4.3, 0, presentation.slide_width / 914400 - 4.3,
             7.5, "203F50", "335C6D")
    nav(slide, values[0])
    title = values[1].replace("，", "，\n", 1).replace("用药动作", "\n用药动作")
    heading = text(slide, title, .65, 1.15, 3.02, 2.85, 30, bold=True)
    heading.text_frame.paragraphs[0].font.size = Pt(23)
    heading.text_frame.paragraphs[0].font.bold = False
    heading.text_frame.paragraphs[0].font.color.rgb = rgb("50646B")
    box(slide, .65, 4.83, .35, .028, THEME["accent"])
    text(slide, values[2].replace("，", "，\n", 1), .65, 5.26,
         3.02, 1.16, 21)
    actions = values[4:]
    for index, (x, y) in enumerate(((4.86, 1.25), (9.06, 1.25),
                                  (4.86, 4.35), (9.06, 4.35))):
        text(slide, actions[index * 2], x, y, 1.49, 1.04, 54,
             color="EEB9AF")
        text(slide, actions[index * 2 + 1], x, y + 1.42, 3.28, .86,
             22, color=THEME["paper"])
    box(slide, 4.86, 3.86, 7.48, .01, "5A7A88")
    box(slide, 8.52, 1.31, .01, 5.51, "5A7A88")
    text(slide, values[3], 11.68, 7.18, 1.0, .2, 11,
         color=THEME["paper"], align=PP_ALIGN.RIGHT)
    return slide


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    if not PREVIOUS.is_file():
        subprocess.run([sys.executable, str(ROOT / "source/build_editorial_texture_samples.py")],
                       check=True, stdout=subprocess.DEVNULL)
    watched = [ORIGINAL, PREVIOUS, ROOT / "交付文件/合理用药_双人逐页讲稿.docx"]
    initial = {str(p.relative_to(ROOT)): hashlib.sha256(p.read_bytes()).hexdigest()
               for p in watched}
    original = Presentation(ORIGINAL)
    result = Presentation()
    result.slide_width, result.slide_height = original.slide_width, original.slide_height
    for page, layout in zip(PAGES, (scene, data, framework)):
        layout(result, original.slides[page - 1])
    verify(original, result)
    samples = OUT / "合理用药_精致编辑版_三页样页.pptx"
    review = OUT / "合理用药_精修样页_对比评审.pptx"
    result.save(samples)
    comparison = compare(original, result, revision_label="精致编辑版")
    prior = compare(Presentation(PREVIOUS), result, baseline_pages=(1, 2, 3),
                    baseline_label="上一版", revision_label="精致编辑版")
    for slide in prior.slides:
        target = comparison.slides.add_slide(comparison.slide_layouts[6])
        clone_shapes(slide, target)
        target.notes_slide.notes_text_frame.text = slide.notes_slide.notes_text_frame.text
    assert len(comparison.slides) == 6
    comparison.save(review)
    for p in watched:
        assert hashlib.sha256(p.read_bytes()).hexdigest() == initial[str(p.relative_to(ROOT))]
    report = {"pages": PAGES, "original_content_notes_and_story_pixels_preserved": True,
              "unchanged_input_hashes": initial,
              "comparison_pages": "1–3 original/refined; 4–6 previous/refined",
              "data_grid": [100, 28], "data_ed": [5, 1],
              "paper_wash_opacity": .64, "samples": str(samples.relative_to(ROOT)),
              "review": str(review.relative_to(ROOT))}
    (OUT / "checks.json").write_text(json.dumps(report, ensure_ascii=False, indent=2))
    print(json.dumps(report, ensure_ascii=False))


if __name__ == "__main__":
    main()
