"""Build three editable review pages without changing either original deliverable.

Run from any directory. Samples stay under work/ until the deck direction is approved.
The approved full-deck generator can reuse the theme and layout helpers below.
"""
from __future__ import annotations

import hashlib
import io
import json
from collections import Counter
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
ORIGINAL = ROOT / "交付文件/合理用药_双人35分钟.pptx"
OUT = ROOT / "work/global-style-samples"
PAGES = (5, 11, 17)
THEME = {
    "font": "Noto Sans CJK SC",
    "ink": "173C70",
    "accent": "F45168",
    "paper": "FFFFFF",
    "tint": "F0F4F9",
    "muted": "D5DFED",
    "rule": "DCE4EE",
    "margin": .65,
    "nav_size": 14,
    "title_size": 40,
    "body_size": 24,
}


def rgb(value):
    return RGBColor.from_string(value)


def box(slide, x, y, w, h, fill, rounded=False):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    obj = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    obj.fill.solid()
    obj.fill.fore_color.rgb = rgb(fill)
    obj.line.fill.background()
    obj._element.spPr.append(OxmlElement("a:effectLst"))
    effect_ref = obj._element.find(".//" + qn("a:effectRef"))
    if effect_ref is not None:
        effect_ref.set("idx", "0")
    if rounded:
        obj.adjustments[0] = .06
    return obj


def text(slide, value, x, y, w, h, size=24, color=None, bold=False,
         align=PP_ALIGN.LEFT):
    obj = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = obj.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    value = value.replace("\\n", "\n")
    for i, part in enumerate(value.split("\n")):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = part
        p.alignment = align
        p.space_before = p.space_after = Pt(0)
        p.line_spacing = 1.15
        p.font.name = THEME["font"]
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = rgb(color or THEME["ink"])
    return obj


def original_texts(slide):
    return [s.text for s in slide.shapes if s.has_text_frame and s.text.strip()]


def canonical(value):
    return "".join(value.replace("\\n", "\n").split())


def base(presentation, original, page):
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(THEME["paper"])
    values = original_texts(original)
    box(slide, .65, .39, .08, .22, THEME["accent"])
    text(slide, values[0], .9, .32, 5.0, .35, 14, bold=True)
    box(slide, .65, .82, 12.03, .012, THEME["rule"])
    title_size = 36 if page == 11 else THEME["title_size"]
    text(slide, values[1], .65, 1.08, 12.03, .76, title_size, bold=True)
    box(slide, .65, 6.57, 12.03, .49, THEME["tint"], rounded=True)
    text(slide, values[2], .87, 6.61, 11.56, .39, 23, bold=True)
    text(slide, values[3], 11.67, 7.17, 1.02, .21, 11, align=PP_ALIGN.RIGHT)
    slide.notes_slide.notes_text_frame.text = original.notes_slide.notes_text_frame.text
    return slide, values[4:]


def picture(slide, original_picture, x, y, w, h):
    """Reuse original pixels and panel crop, contain without stretching or new crops."""
    image = original_picture.image
    l, t, r, b = (getattr(original_picture, "crop_" + edge)
                  for edge in ("left", "top", "right", "bottom"))
    iw, ih = image.size
    ratio = iw * (1 - l - r) / (ih * (1 - t - b))
    if w / h > ratio:
        actual_w, actual_h = h * ratio, h
    else:
        actual_w, actual_h = w, w / ratio
    obj = slide.shapes.add_picture(io.BytesIO(image.blob), Inches(x), Inches(y),
                                   width=Inches(actual_w), height=Inches(actual_h))
    obj.crop_left, obj.crop_top, obj.crop_right, obj.crop_bottom = l, t, r, b
    obj.name = original_picture.name
    return obj


def scene(presentation, original):
    slide, captions = base(presentation, original, 5)
    panels = [s for s in original.shapes if s.shape_type == 13]
    assert len(panels) == 3 and len(captions) == 3
    picture(slide, panels[0], .65, 2.13, 4.82, 3.2)
    box(slide, .65, 5.49, 4.82, .75, THEME["tint"], rounded=True)
    text(slide, captions[0], .85, 5.58, 4.42, .53, 23, bold=True)
    picture(slide, panels[1], 5.91, 2.13, 3.18, 2.11)
    box(slide, 9.36, 2.3, 3.32, 1.4, THEME["tint"], rounded=True)
    text(slide, captions[1].replace("，", "，\n", 1), 9.58, 2.5, 2.88, 1.0, 23, bold=True)
    picture(slide, panels[2], 5.91, 4.34, 3.18, 2.11)
    box(slide, 9.36, 4.66, 3.32, 1.4, THEME["tint"], rounded=True)
    text(slide, captions[2].replace("，", "，\n", 1), 9.58, 4.86, 2.88, 1.0, 23, bold=True)
    return slide


def emergency_icon(slide, x, y, w, color, name):
    # Counting symbols, not medical anatomy. All instances have identical geometry.
    obj = box(slide, x, y, w, 1.05, color, rounded=True)
    obj.name = name
    box(slide, x + w*.4, y + .16, w*.2, .28, THEME["paper"])
    box(slide, x + w*.27, y + .24, w*.46, .12, THEME["paper"])
    box(slide, x + w*.43, y + .65, w*.18, .4, THEME["paper"])
    box(slide, x + w*.16, y + .56, w*.12, .12, THEME["paper"])
    box(slide, x + w*.73, y + .56, w*.12, .12, THEME["paper"])


def data(presentation, original):
    slide, values = base(presentation, original, 11)
    assert len(values) == 8
    text(slide, values[0], .65, 2.2, 6.96, .42, 23, bold=True)
    text(slide, values[1], .65, 3.04, 2.86, .36, 23,
         color=THEME["accent"], bold=True)
    text(slide, values[2], .62, 3.54, 3.18, 1.25, 84,
         color=THEME["accent"], bold=True)
    text(slide, values[3], .65, 5.62, 6.98, .45, 27, bold=True)
    for index in range(100):
        row, col = divmod(index, 10)
        obj = box(slide, 4.49 + col * .302, 2.94 + row * .302,
                  .241, .241, THEME["accent"] if index < 28 else THEME["muted"],
                  rounded=True)
        obj.name = f"DATA_GRID_{index+1:03d}_{'HIGHLIGHT' if index < 28 else 'REST'}"
    box(slide, 8.1, 2.08, 4.58, 4.23, THEME["tint"], rounded=True)
    text(slide, values[4], 8.35, 2.33, 4.08, .48, 22, bold=True)
    text(slide, values[5], 8.35, 3.04, 4.08, .62, 33,
         color=THEME["accent"], bold=True)
    text(slide, values[6], 8.35, 3.82, 4.08, .4, 23, bold=True)
    for i in range(5):
        emergency_icon(slide, 8.36 + i * .815, 4.55, .65,
                       THEME["accent"] if i == 0 else THEME["ink"],
                       f"DATA_ED_{i+1}_{'HIGHLIGHT' if i == 0 else 'REST'}")
    text(slide, values[7], 8.35, 5.85, 4.08, .32, 18, bold=True)
    return slide


def framework(presentation, original):
    slide, values = base(presentation, original, 17)
    assert len(values) == 8
    for index in range(4):
        y = 2.1 + index * 1.085
        text(slide, values[index * 2], .84, y, 1.17, .88, 50,
             color=THEME["accent"], bold=True, align=PP_ALIGN.CENTER)
        box(slide, 2.45, y, 10.23, .89, THEME["tint"], rounded=True)
        description = values[index * 2 + 1].replace("\\n", "").replace("\n", "")
        text(slide, description, 2.73, y + .15, 9.67, .58, 27, bold=True)
    return slide


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    before = hashlib.sha256(ORIGINAL.read_bytes()).hexdigest()
    original = Presentation(ORIGINAL)
    result = Presentation()
    result.slide_width, result.slide_height = original.slide_width, original.slide_height
    for page, layout in zip(PAGES, (scene, data, framework)):
        fresh = layout(result, original.slides[page - 1])
        expected = Counter(canonical(t) for t in original_texts(original.slides[page - 1]))
        actual = Counter(canonical(t) for t in original_texts(fresh))
        assert actual == expected, {"page": page, "missing": expected - actual,
                                    "added": actual - expected}
        assert fresh.notes_slide.notes_text_frame.text == original.slides[page - 1].notes_slide.notes_text_frame.text
        original_images = Counter((s.name, hashlib.sha256(s.image.blob).hexdigest())
                                  for s in original.slides[page - 1].shapes if s.shape_type == 13)
        fresh_images = Counter((s.name, hashlib.sha256(s.image.blob).hexdigest())
                               for s in fresh.shapes if s.shape_type == 13)
        assert original_images == fresh_images
        for obj in fresh.shapes:
            assert obj.left >= 0 and obj.top >= 0
            assert obj.left + obj.width <= result.slide_width
            assert obj.top + obj.height <= result.slide_height
    shapes = result.slides[1].shapes
    assert sum(s.name.startswith("DATA_GRID_") for s in shapes) == 100
    assert sum(s.name.startswith("DATA_GRID_") and s.name.endswith("HIGHLIGHT")
               for s in shapes) == 28
    assert sum(s.name.startswith("DATA_ED_") for s in shapes) == 5
    assert sum(s.name.startswith("DATA_ED_") and s.name.endswith("HIGHLIGHT")
               for s in shapes) == 1
    path = OUT / "合理用药_全局样式_样页.pptx"
    result.save(path)
    assert hashlib.sha256(ORIGINAL.read_bytes()).hexdigest() == before
    report = {"pages": PAGES, "theme": THEME, "text_and_notes_preserved": True,
              "original_unchanged": True, "original_sha256": before,
              "data_grid": {"total": 100, "highlight": 28},
              "data_ed": {"total": 5, "highlight": 1},
              "sample": str(path.relative_to(ROOT))}
    (OUT / "review-checks.json").write_text(json.dumps(report, ensure_ascii=False, indent=2))
    print(json.dumps(report, ensure_ascii=False))


if __name__ == "__main__":
    main()
